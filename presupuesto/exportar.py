"""Exporta el resultado a Excel: resumen, líneas con su cálculo y revisión.

Es un archivo de trabajo para copiar al formato de cada donante, no el formato
del donante en sí.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calculo import Contexto, Resultado
from .modelo import NOMBRE_RUBRO, Banco, Proyecto

AZUL = "1F4E5A"
H = Font(bold=True, color="FFFFFF")
HF = PatternFill("solid", fgColor=AZUL)
ROJO = PatternFill("solid", fgColor="F4B6B6")
AMAR = PatternFill("solid", fgColor="FFE699")
VERDE = PatternFill("solid", fgColor="C6EFCE")


def _tabla(ws, fila, cabeceras, anchos):
    for i, (c, w) in enumerate(zip(cabeceras, anchos), 1):
        x = ws.cell(row=fila, column=i, value=c)
        x.font, x.fill = H, HF
        x.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def exportar(banco: Banco, proyecto: Proyecto, res: Resultado, ruta: Path) -> Path:
    par = proyecto.parametros
    mon, fx = par.moneda, res.fx
    fmt = '#,##0.00' if mon != "COP" else '#,##0'
    cx = Contexto(banco, proyecto)
    nombres = {a.codigo: a.nombre for a in proyecto.actividades}
    wb = Workbook()

    # ------------------------------------------------------------------ Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = par.nombre
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws["A2"] = f"Donante: {par.donante} · Moneda: {mon} · TRM USD {par.trm_usd:,.0f} · Indirectos {par.pct_indirectos:.1%}"
    ws["A3"] = f"Revisión: {res.rojos()} rojos · {res.amarillos()} amarillos"
    ws["A3"].fill = ROJO if res.rojos() else VERDE
    _tabla(ws, 5, ["Categoría del donante", f"Total {mon}", "% del directo"], [60, 18, 14])
    orden = banco.categorias.get(par.donante, [])
    cats = sorted(res.por_categoria, key=lambda c: (orden.index(c) if c in orden else 99, c))
    r = 6
    for c in cats:
        ws.cell(row=r, column=1, value=c)
        ws.cell(row=r, column=2, value=res.por_categoria[c] / fx).number_format = fmt
        ws.cell(row=r, column=3, value=res.por_categoria[c] / res.directos_cop if res.directos_cop else 0).number_format = "0.0%"
        r += 1
    for lab, v in [("TOTAL COSTOS DIRECTOS", res.directos_cop), (f"Costos indirectos ({par.pct_indirectos:.1%})", res.indirectos_cop),
                   ("TOTAL PROYECTO", res.total_cop)]:
        ws.cell(row=r, column=1, value=lab).font = Font(bold=True)
        x = ws.cell(row=r, column=2, value=v / fx)
        x.number_format, x.font = fmt, Font(bold=True)
        r += 1
    r += 2
    ws.cell(row=r, column=1, value="Por actividad").font = Font(bold=True, color=AZUL, size=12)
    r += 1
    for i, h in enumerate(["Código", "Actividad", f"Total {mon}"], 1):
        x = ws.cell(row=r, column=i, value=h)
        x.font, x.fill = H, HF
    ws.column_dimensions["B"].width = 18
    por_act: dict[str, float] = {}
    for l in res.lineas:
        por_act[l.actividad] = por_act.get(l.actividad, 0) + l.total_cop
    for cod in sorted(por_act, key=lambda c: [int(p) if p.isdigit() else p for p in c.split(".")]):
        r += 1
        ws.cell(row=r, column=1, value=cod)
        ws.cell(row=r, column=2, value=nombres.get(cod, ""))
        ws.cell(row=r, column=3, value=por_act[cod] / fx).number_format = fmt

    # ------------------------------------------------------------------ Líneas
    wl = wb.create_sheet("Líneas")
    cab = ["Código", "Actividad", "Lugar", "Grupo / cargo / concepto", "Rubro", "Bloque interno", f"Categoría {par.donante}",
           "Cantidad", "Unidad", f"Costo unitario {mon}", f"Total {mon}", "Total COP", "Cálculo (para la narrativa)"]
    _tabla(wl, 1, cab, [9, 36, 22, 30, 20, 26, 36, 10, 18, 14, 16, 16, 80])
    orden_l = sorted(res.lineas, key=lambda l: ([int(p) if p.isdigit() else 0 for p in l.actividad.split(".")],
                                                {"personal": 0, "evento": 1, "otro": 2}[l.origen], l.evento or 0))
    for i, l in enumerate(orden_l, 2):
        cat = cx.mapeo.get(l.bloque, {}).get(par.donante, "SIN MAPEO")
        unit = l.unitario_cop / fx
        vals = [l.actividad, nombres.get(l.actividad, ""), l.lugar, l.detalle, l.rubro, l.bloque, cat, l.cantidad, l.unidad,
                unit, l.total_cop / fx, l.total_cop,
                f"{l.calculo} × {unit:,.2f} {mon} = {l.total_cop / fx:,.2f} {mon}"]
        for j, v in enumerate(vals, 1):
            x = wl.cell(row=i, column=j, value=v)
            if j in (10, 11):
                x.number_format = fmt
            elif j == 12:
                x.number_format = '#,##0'
    wl.auto_filter.ref = f"A1:{get_column_letter(len(cab))}{len(orden_l) + 1}"

    # ------------------------------------------------------------------ Revisión
    wr = wb.create_sheet("Revisión")
    _tabla(wr, 1, ["Nivel", "Dónde", "Qué revisar"], [10, 60, 100])
    i = 2
    for a in sorted(res.avisos, key=lambda a: a.nivel != "rojo"):
        wr.cell(row=i, column=1, value=a.nivel.upper()).fill = ROJO if a.nivel == "rojo" else AMAR
        wr.cell(row=i, column=2, value=a.donde)
        wr.cell(row=i, column=3, value=a.texto).alignment = Alignment(wrap_text=True)
        i += 1
    i += 1
    wr.cell(row=i, column=1, value="Rubros marcados como no necesarios").font = Font(bold=True, color=AZUL)
    i += 1
    for c in res.celdas:
        if c.estado == "no_necesario":
            ev = proyecto.eventos[c.evento]
            wr.cell(row=i, column=1, value="—")
            wr.cell(row=i, column=2, value=f"{ev.actividad} · {ev.lugar} · {c.grupo or 'Evento'} · {NOMBRE_RUBRO[c.rubro]}")
            wr.cell(row=i, column=3, value=c.razon)
            i += 1
    for it, estado, razon in res.chequeo:
        if estado == "no_necesario":
            wr.cell(row=i, column=1, value="—")
            wr.cell(row=i, column=2, value=f"Lista de chequeo · {it.nombre}")
            wr.cell(row=i, column=3, value=razon)
            i += 1

    ruta = Path(ruta)
    ruta.parent.mkdir(parents=True, exist_ok=True)
    wb.save(ruta)
    return ruta
