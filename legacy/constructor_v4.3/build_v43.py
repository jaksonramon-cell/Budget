import sys, json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo, TableFormula
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.formatting.rule import CellIsRule, FormulaRule, Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
from openpyxl.comments import Comment
from openpyxl.workbook.properties import CalcProperties

DATA = sys.argv[1]
OUT = sys.argv[2]
V3 = json.load(open(DATA))
PROJ = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}

wb = Workbook()
FONT = "Arial"
HDR = PatternFill("solid", fgColor="1F4E5A")
HDRC = PatternFill("solid", fgColor="7F7F7F")
HDRA = PatternFill("solid", fgColor="B07D2B")   # header of override (ajuste) columns
INP = PatternFill("solid", fgColor="FFF8DC")
OPT = PatternFill("solid", fgColor="FFFDF2")    # optional inputs
CAL = PatternFill("solid", fgColor="F2F2F2")
QTY = PatternFill("solid", fgColor="E8EEF4")
TILE = PatternFill("solid", fgColor="EAF1F4")
TITLE = Font(name=FONT, size=14, bold=True, color="1F4E5A")
H2 = Font(name=FONT, size=11, bold=True, color="1F4E5A")
SUB = Font(name=FONT, size=9, italic=True, color="595959")
HF = Font(name=FONT, size=9, bold=True, color="FFFFFF")
BF = Font(name=FONT, size=9)
BLUE = Font(name=FONT, size=9, color="0000FF")
BOLD = Font(name=FONT, size=10, bold=True)
thin = Side(style="thin", color="D9D9D9")
BRD = Border(left=thin, right=thin, top=thin, bottom=thin)
COP = '#,##0;(#,##0);"-"'
USD = '#,##0.00;(#,##0.00);"-"'
PRES = COP
NUM = '#,##0;(#,##0);"-"'
PCT = '0%'
TXT = '@'
RED = PatternFill("solid", fgColor="F4B6B6"); YEL = PatternFill("solid", fgColor="FFE699")
GRN = PatternFill("solid", fgColor="C6EFCE"); ORG = PatternFill("solid", fgColor="F8CBAD")
GROW = 1000
MAX_ACT_TAB = 60      # activities shown in Tablero
MAX_ACT_LD = 100      # activities in LineasDonante
MAX_DEP = 12


def name(n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def title(ws, t, s):
    ws["A1"] = t; ws["A1"].font = TITLE
    ws["A2"] = s; ws["A2"].font = SUB
    ws.sheet_view.showGridLines = False


def dv_list(ws, formula, rng):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Elija un valor de la lista"; dv.showErrorMessage = True
    ws.add_data_validation(dv); dv.add(rng)


def estado_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Confirmado"'], fill=GRN))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Estimado"'], fill=YEL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Por confirmar"'], fill=ORG))


def money_cf(ws, rng):
    dxf = DifferentialStyle(numFmt=NumberFormat(numFmtId=200, formatCode='#,##0.00;(#,##0.00);"-"'))
    ws.conditional_formatting.add(rng, Rule(type="expression", dxf=dxf, formula=['Moneda="USD"']))


def mtxt(x):
    return f'FIXED({x},Dec)&" "&Moneda'


def hcell(ws, r, c, v, fill=HDR):
    x = ws.cell(row=r, column=c, value=v); x.font = HF; x.fill = fill; x.border = BRD
    x.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    return x


def bcell(ws, r, c, v, fmt=None, fill=CAL, font=BF):
    x = ws.cell(row=r, column=c, value=v); x.font = font; x.fill = fill; x.border = BRD
    if fmt: x.number_format = fmt
    return x


class Tbl:
    """Excel Table. cols: dicts {n, f (calc formula), w, fmt, kind: qty|opt|adj, note, wrap}"""
    def __init__(self, ws, tname, cols, rows, hrow=4, min_rows=1):
        self.ws, self.t, self.cols, self.h = ws, tname, cols, hrow
        self.idx = {c["n"]: i + 1 for i, c in enumerate(cols)}
        n = max(len(rows), min_rows)
        self.r1, self.r2 = hrow + 1, hrow + n
        for i, c in enumerate(cols, 1):
            calc = c.get("f") is not None
            fill = HDRC if calc else (HDRA if c.get("kind") == "adj" else HDR)
            cell = hcell(ws, hrow, i, c["n"], fill)
            ws.column_dimensions[L(i)].width = c.get("w", 12)
            if c.get("note"):
                cell.comment = Comment(c["note"], "CIREC"); cell.comment.width = 300; cell.comment.height = 120
        ws.row_dimensions[hrow].height = 52
        for k in range(n):
            r = hrow + 1 + k
            row = rows[k] if k < len(rows) else {}
            for i, c in enumerate(cols, 1):
                cell = ws.cell(row=r, column=i)
                if c.get("f") is not None:
                    cell.value = "=" + c["f"]
                    cell.fill = QTY if c.get("kind") == "qty" else CAL
                else:
                    v = row.get(c["n"])
                    if v is not None:
                        cell.value = v
                    cell.fill = OPT if c.get("kind") in ("opt", "adj") else INP
                cell.font = BF; cell.border = BRD
                if c.get("fmt"): cell.number_format = c["fmt"]
                if c.get("wrap"): cell.alignment = Alignment(wrap_text=True, vertical="top")
        t = Table(displayName=tname, ref=f"A{hrow}:{L(len(cols))}{self.r2}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        t._initialise_columns()
        for i, c in enumerate(cols):
            t.tableColumns[i].name = c["n"]
            if c.get("f") is not None:
                t.tableColumns[i].calculatedColumnFormula = TableFormula(attr_text=c["f"])
        ws.add_table(t)
        ws.freeze_panes = ws.cell(row=hrow + 1, column=1)

    def col(self, n):
        return L(self.idx[n])

    def rng(self, n):
        c = self.col(n)
        return f"{c}{self.r1}:{c}{self.h + GROW}"


def R(t, c):
    return f"{t}[[#This Row],[{c}]]"


def rate(i, z):
    a = f'SUMIFS(tblTarifas[Valor COP],tblTarifas[ID tarifa],{i},tblTarifas[Zona],{z})'
    b = f'SUMIFS(tblTarifas[Valor COP],tblTarifas[ID tarifa],{i},tblTarifas[Zona],"Todas")'
    return f'IF({a}>0,{a},{b})'


def rows_of(tn):
    return V3[tn]["rows"]


def dt(v):
    if isinstance(v, str) and len(v) == 10 and v[4] == "-":
        return datetime.date.fromisoformat(v)
    return v


# ================================================================= LISTAS (hidden)
wsL = wb.active; wsL.title = "Listas"
DEPS = ["Amazonas", "Antioquia", "Arauca", "Atlántico", "Bogotá D.C.", "Bolívar", "Boyacá",
        "Caldas", "Caquetá", "Casanare", "Cauca", "Cesar", "Chocó", "Córdoba", "Cundinamarca", "Guainía",
        "Guaviare", "Huila", "La Guajira", "Magdalena", "Meta", "Nariño", "Norte de Santander", "Putumayo",
        "Quindío", "Risaralda", "San Andrés", "Santander", "Sucre", "Tolima", "Valle del Cauca", "Vaupés", "Vichada"]
lists = {
    "A": ("Zonas", ["Todas", "Bogotá", "Capital departamental", "Municipio", "Rural remoto"]),
    "B": ("Estados", ["Confirmado", "Estimado", "Por confirmar"]),
    "C": ("Alimentacion", ["Completa", "Refrigerio", "Ninguna"]),
    "D": ("Niveles", ["Resultado", "Producto", "Actividad"]),
    "E": ("Departamentos", DEPS),
    "F": ("Monedas", ["COP", "USD"]),
    "G": ("SiNo", ["Sí", "No"]),
    "H": ("Si", ["Sí"]),
    "I": ("BaseInd", ["Costos directos", "Total del proyecto"]),
}
wsL["A1"] = "Listas de los menús desplegables"; wsL["A1"].font = TITLE
for col, (h, vals) in lists.items():
    wsL[f"{col}3"] = h; wsL[f"{col}3"].font = HF; wsL[f"{col}3"].fill = HDR
    for i, v in enumerate(vals):
        wsL[f"{col}{4+i}"] = v; wsL[f"{col}{4+i}"].font = BF
    wsL.column_dimensions[col].width = 24
    name("Lista_" + h, f"Listas!${col}$4:${col}${3+len(vals)}")
name("DV_WBS", "tblWBS[Código WBS]")
name("DV_Cargos", "tblCargos[Cargo]")
name("DV_Tarifas", "tblTarifas[ID tarifa]")
name("DV_Bloques", "tblMapeo[Bloque]")
name("DV_Plantillas", "tblPlantillas[Plantilla]")
name("DV_Lugares", "tblLugares[Lugar]")

# ================================================================= PARAMETROS
wsP = wb.create_sheet("Parametros")
title(wsP, "Parámetros del proyecto", "Celdas amarillas = editar. Internamente todo se calcula en COP.")
for i, (h, w) in enumerate([("Parámetro", 58), ("Valor", 40), ("Nota", 80)], 1):
    hcell(wsP, 4, i, h); wsP.column_dimensions[L(i)].width = w
pv = dict(V3["params"]); pv.update(PROJ.get("params", {}))
params = [
    ("Proyecto", pv.get("Proyecto"), None, None, "Proyecto"),
    ("Donante activo (define las categorías del Resumen)", pv.get("Donante activo (define las categorías del Resumen)"), None, "Los donantes disponibles son los encabezados de la hoja MapeoDonante", "DonanteActivo"),
    ("Moneda de presentación", pv.get("Moneda de presentación"), None, "COP o USD. Cambia totales y narrativas. No afecta los cálculos internos ni Power BI.", "Moneda"),
    ("TRM (COP por USD)", pv.get("TRM (COP por USD)"), COP, "TRM de formulación del proyecto", "TRM"),
    ("% costos indirectos", pv.get("% costos indirectos"), PCT, None, "PctInd"),
    ("Base de los indirectos", pv.get("Base de los indirectos"), None, "Costos directos: indirectos = % × directos. Total del proyecto: el indirecto es ese % del total final (directos ÷ (1 − %) × %). Confirmar con el donante.", "BaseInd"),
    ("Máximo % de indirectos permitido por el donante", pv.get("Máximo % de indirectos permitido por el donante"), PCT, None, "MaxInd"),
    ("Techo presupuestal del donante (en moneda de presentación)", pv.get("Techo"), COP, "Dejar vacío si no aplica", "Techo"),
    ("Total del formato del donante (en moneda de presentación)", pv.get("TotalDonante"), COP, "Pegar aquí el total del Excel del donante antes de enviar", "TotalDonante"),
]
for i, (lab, val, fmt, note, nm) in enumerate(params):
    r = 5 + i
    wsP.cell(row=r, column=1, value=lab).font = BF
    c = wsP.cell(row=r, column=2, value=val); c.fill = INP; c.font = BLUE; c.border = BRD
    if fmt: c.number_format = fmt
    wsP.cell(row=r, column=3, value=note).font = SUB
    wsP.cell(row=r, column=3).alignment = Alignment(wrap_text=True)
    name(nm, f"Parametros!$B${r}")
wsP["B7"].font = Font(name=FONT, size=11, bold=True, color="0000FF")
r = 5 + len(params) + 1
for lab, f, nm in [("(interno) índice del donante activo", "=MATCH(DonanteActivo,MapeoDonante!$B$4:$E$4,0)", "DonIdx"),
                   ("(interno) factor a moneda de presentación", '=IF(Moneda="USD",1/TRM,1)', "FX"),
                   ("(interno) decimales de presentación", '=IF(Moneda="USD",2,0)', "Dec")]:
    wsP.cell(row=r, column=1, value=lab).font = SUB
    wsP.cell(row=r, column=2, value=f).font = SUB
    name(nm, f"Parametros!$B${r}"); r += 1
dv_list(wsP, "=MapeoDonante!$B$4:$E$4", "B6")
dv_list(wsP, "=Lista_Monedas", "B7")
dv_list(wsP, "=Lista_BaseInd", "B10")
money_cf(wsP, "B12:B13")

# ================================================================= TARIFAS (+ Bloque)
wsT = wb.create_sheet("Tarifas")
title(wsT, "Banco de precios – una sola tabla para todos los precios",
      "Un ID puede tener una fila por zona; si no hay para la zona del evento se usa 'Todas'. 'Bloque' = a dónde va el costo cuando se usa en OtrosCostos. Agregar: escriba en la fila justo debajo de la tabla.")
T = "tblTarifas"
tcols = [
    dict(n="ID tarifa", w=14, fmt=TXT), dict(n="Categoría", w=14), dict(n="Concepto", w=46), dict(n="Unidad", w=14),
    dict(n="Bloque", w=24, note="Bloque de costo que se asigna automáticamente cuando esta tarifa se usa en OtrosCostos."),
    dict(n="Zona", w=18), dict(n="Valor", w=13, fmt='#,##0.##', note="En la moneda de la cotización (columna Moneda)."),
    dict(n="Moneda", w=8),
    dict(n="Valor COP", w=13, fmt=COP, f=f'IF({R(T,"Valor")}="","",IF({R(T,"Moneda")}="USD",{R(T,"Valor")}*TRM,{R(T,"Valor")}))'),
    dict(n="Valor USD", w=11, fmt=USD, f=f'IF({R(T,"Valor COP")}="","",{R(T,"Valor COP")}/TRM)'),
    dict(n="En uso", w=8, f=(f'IF({R(T,"ID tarifa")}="","",IF(COUNTIF(tblOtros[ID tarifa],{R(T,"ID tarifa")})+COUNTIF(tblEventos[Tarifa transporte participantes],{R(T,"ID tarifa")})'
                             f'+COUNTIF(tblEventos[Kit por participante],{R(T,"ID tarifa")})+COUNTIF(tblEventos[Paquete por evento],{R(T,"ID tarifa")})+ISNUMBER(MATCH({R(T,"ID tarifa")},{{"HOSP-STAFF","HOSP-PART","ALIM","REFRIG","TIQ","TTE-STAFF","SALON"}},0))>0,"Sí",""))'),
         note="Sí = este proyecto usa la tarifa. Filtre por Sí para revisar solo los precios que importan."),
    dict(n="Fuente", w=34), dict(n="Fecha", w=11, fmt="dd/mm/yyyy"), dict(n="Estado", w=13), dict(n="Notas", w=40),
]
BLOQUE_POR_ID = {
    "TALLER-ERAE": "Paquetes por evento",
    "HOSP-STAFF": "Hospedaje staff", "HOSP-PART": "Hospedaje participantes", "ALIM": "Alimentación participantes",
    "REFRIG": "Alimentación participantes", "TIQ": "Tiquetes aéreos participantes", "TTE-STAFF": "Transporte terrestre staff",
    "TTE-PART-LOCAL": "Transporte participantes", "TTE-PART-MUNIC": "Transporte participantes", "SALON": "Alquiler de salones",
    "KIT-PED": "Materiales y kits", "KIT-ENT": "Materiales y kits", "KIT-ERAE": "Materiales y kits", "MAT-DIF": "Materiales y kits",
    "KIT-SENAL": "Materiales y kits", "DOTAC": "Materiales y kits", "DISP-MOV": "Dispositivos de movilidad",
    "CASH-EMERG": "Transferencias monetarias", "SEG-VIDA": "Seguros", "CEL-MES": "Gastos operativos", "LAPTOP": "Equipos",
}
trows = []
for t in rows_of("tblTarifas"):
    d = dict(t)
    d["Bloque"] = BLOQUE_POR_ID.get(t["ID tarifa"], "Dispositivos de movilidad" if t.get("Categoría") == "Prótesis y órtesis" else None)
    d["Fecha"] = dt(t.get("Fecha"))
    trows.append(d)
for ov in PROJ.get("tarifas", []):
    hit = [t for t in trows if t["ID tarifa"] == ov["ID tarifa"] and t["Zona"] == ov["Zona"]]
    if hit:
        hit[0].update(ov)
    else:
        base = next((t for t in trows if t["ID tarifa"] == ov["ID tarifa"]), {})
        nd = {k: base.get(k) for k in ("Categoría", "Concepto", "Unidad", "Bloque")}; nd.update(ov)
        idx = max(i for i, t in enumerate(trows) if t["ID tarifa"] == ov["ID tarifa"]) if base else len(trows) - 1
        trows.insert(idx + 1, nd)
trows.append({"ID tarifa": "TALLER-ERAE", "Categoría": "Talleres", "Concepto": "Taller comunitario ERAE / MCMT todo incluido (salón, refrigerio, materiales)",
              "Unidad": "taller", "Bloque": "Paquetes por evento", "Zona": "Todas", "Valor": 300000, "Moneda": "COP",
              "Fuente": "EJEMPLO – reemplazar por cotización", "Estado": "Por confirmar", "Notas": "Valor ilustrativo, no es un precio real."})
TT = Tbl(wsT, T, tcols, trows)
dv_list(wsT, "=Lista_Zonas", TT.rng("Zona")); dv_list(wsT, "=Lista_Monedas", TT.rng("Moneda"))
dv_list(wsT, "=Lista_Estados", TT.rng("Estado")); dv_list(wsT, "=DV_Bloques", TT.rng("Bloque"))
estado_cf(wsT, TT.rng("Estado"))

# ================================================================= CARGOS
wsC = wb.create_sheet("Cargos")
title(wsC, "Costo de personal por cargo (costo total mensual para CIREC)", "Por cargo, no por persona: salario + auxilio de transporte + carga prestacional.")
CG = "tblCargos"
ccols = [
    dict(n="Cargo", w=44), dict(n="Salario base COP", w=15, fmt=COP), dict(n="Costo mensual COP", w=16, fmt=COP),
    dict(n="Costo día COP", w=14, fmt=COP, f=f'IF({R(CG,"Costo mensual COP")}="","",{R(CG,"Costo mensual COP")}/30)'),
    dict(n="Costo mensual USD", w=14, fmt=USD, f=f'IF({R(CG,"Costo mensual COP")}="","",{R(CG,"Costo mensual COP")}/TRM)'),
    dict(n="Fuente", w=34), dict(n="Estado", w=14), dict(n="Notas", w=50),
]
TC = Tbl(wsC, CG, ccols, PROJ.get("cargos", rows_of("tblCargos")))
dv_list(wsC, "=Lista_Estados", TC.rng("Estado")); estado_cf(wsC, TC.rng("Estado"))

# ================================================================= LUGARES (new)
wsG = wb.create_sheet("Lugares")
title(wsG, "Lugares – cada municipio se registra UNA vez con su departamento y su zona tarifaria",
      "En Eventos solo se elige el lugar: el departamento y la zona salen de aquí. Para un municipio aún no definido use 'Por definir (Departamento)'.")
CAPS = [("Leticia", "Amazonas"), ("Medellín", "Antioquia"), ("Arauca", "Arauca"), ("Barranquilla", "Atlántico"), ("Cartagena", "Bolívar"),
        ("Tunja", "Boyacá"), ("Manizales", "Caldas"), ("Florencia", "Caquetá"), ("Yopal", "Casanare"), ("Popayán", "Cauca"),
        ("Valledupar", "Cesar"), ("Quibdó", "Chocó"), ("Montería", "Córdoba"), ("Inírida", "Guainía"), ("San José del Guaviare", "Guaviare"),
        ("Neiva", "Huila"), ("Riohacha", "La Guajira"), ("Santa Marta", "Magdalena"), ("Villavicencio", "Meta"), ("Pasto", "Nariño"),
        ("Cúcuta", "Norte de Santander"), ("Mocoa", "Putumayo"), ("Armenia", "Quindío"), ("Pereira", "Risaralda"), ("San Andrés", "San Andrés"),
        ("Bucaramanga", "Santander"), ("Sincelejo", "Sucre"), ("Ibagué", "Tolima"), ("Cali", "Valle del Cauca"), ("Mitú", "Vaupés"),
        ("Puerto Carreño", "Vichada")]
lrows = [{"Lugar": "Bogotá", "Departamento": "Bogotá D.C.", "Zona": "Bogotá", "Notas": "Sede CIREC"}]
lrows += [{"Lugar": m, "Departamento": d, "Zona": "Capital departamental"} for m, d in CAPS]
lrows += [{"Lugar": "Miraflores (Guaviare)", "Departamento": "Guaviare", "Zona": "Rural remoto", "Notas": "Acceso aéreo desde San José (PRESUP_1)"}]
lrows += [{"Lugar": f"Por definir ({d})", "Departamento": d, "Zona": "Municipio", "Notas": "Comodín mientras se define el municipio"}
          for d in ["Antioquia", "Arauca", "Caquetá", "Cauca", "Chocó", "Guaviare", "Nariño", "Norte de Santander", "Putumayo", "Meta"]]
lrows += PROJ.get("lugares_extra", [])
gcols = [dict(n="Lugar", w=30), dict(n="Departamento", w=20), dict(n="Zona", w=22, note="Define qué precio se usa (Tarifas)."), dict(n="Notas", w=50)]
TG = Tbl(wsG, "tblLugares", gcols, lrows)
dv_list(wsG, "=Lista_Departamentos", TG.rng("Departamento")); dv_list(wsG, "=Lista_Zonas", TG.rng("Zona"))

# ================================================================= PLANTILLAS (new)
wsK = wb.create_sheet("Plantillas")
title(wsK, "Plantillas de evento – el perfil típico de cada tipo de actividad (se define UNA vez)",
      "Cada evento hereda estos valores. Si un evento es distinto, se ajusta solo esa celda en Eventos (columnas 'Ajuste'). Cambiar aquí cambia todos los eventos que usan la plantilla.")
PARAMS_NUM = ["Staff viaja", "Staff local", "Staff en avión", "Días staff en sitio", "Noches staff", "Días transporte staff", "Participantes", "Acompañantes",
              "Días participantes", "Noches participantes", "Porcentaje con hospedaje", "Personas por habitación",
              "Porcentaje con transporte", "Días de salón"]
PARAMS_TXT = ["Alimentación participantes", "Tarifa transporte participantes", "Kit por participante", "Paquete por evento", "Mismos participantes en repeticiones"]
NOTES = {
    "Staff viaja": "Personal CIREC que se desplaza: genera hospedaje, alimentación y transporte terrestre.",
    "Staff local": "Personal que ya está en el lugar: solo alimentación.",
    "Staff en avión": "Cuántos del staff que viaja van en avión (1 tiquete ida y regreso c/u).",
    "Días staff en sitio": "CONVENCIÓN: días contando los de viaje. Se usa para alimentación y transporte.",
    "Noches staff": "CONVENCIÓN: noches que duermen fuera. Normalmente días − 1.",
    "Días transporte staff": "Días con transporte terrestre local por persona (aeropuerto + desplazamientos internos). Puede ser menor que los días en sitio.",
    "Participantes": "Participantes por evento (sin acompañantes).",
    "Porcentaje con hospedaje": "100% si todos duermen allí; 40% si solo los que vienen de lejos.",
    "Personas por habitación": "2 = acomodación doble (ej. beneficiario + acompañante).",
    "Tarifa transporte participantes": "ID de tarifa de transporte (ej. TTE-PART-LOCAL).",
    "Kit por participante": "ID de tarifa del kit que recibe cada participante. Vacío si no aplica.",
    "Paquete por evento": "ID de tarifa de un precio TODO INCLUIDO por evento (ej. taller comunitario con salón, refrigerio y materiales). Se cobra 1 por cada repetición. Vacío si no aplica.",
    "Mismos participantes en repeticiones": "Sí = las repeticiones son con las mismas personas (ej. 3 fases de una brigada): se cuentan una vez. No = cada repetición tiene personas nuevas.",
}
PK = "tblPlantillas"
kcols = [dict(n="Plantilla", w=30), dict(n="Descripción", w=36, wrap=True)]
for p in PARAMS_NUM:
    kcols.append(dict(n=p, w=10, fmt=PCT if p.startswith("Porcentaje") else None, note=NOTES.get(p)))
for p in PARAMS_TXT:
    kcols.append(dict(n=p, w=13, note=NOTES.get(p)))
kcols += [dict(n="Estado", w=13), dict(n="Notas", w=44, wrap=True)]
def tpl(name_, desc, sv, sl, sa, ds, ns, p, a, dp, np_, ph, pph, pt, sal, alim, ttp, kit, mismos, est, notes, td=None, pkg=None):
    v = dict(zip([x for x in PARAMS_NUM if x != "Días transporte staff"], [sv, sl, sa, ds, ns, p, a, dp, np_, ph, pph, pt, sal]))
    v["Días transporte staff"] = ds if td is None else td
    v.update(dict(zip(PARAMS_TXT, [alim, ttp, kit, pkg, mismos])))
    v.update({"Plantilla": name_, "Descripción": desc, "Estado": est, "Notas": notes})
    return v
krows = [
    tpl("Brigada de rehabilitación (fase)", "Una fase de brigada en un lugar. Repeticiones = nº de fases (normalmente 3).",
        7, 0, 7, 4, 3, 12, 12, 2, 1, 1, 2, 1, 0, "Completa", "TTE-PART-LOCAL", None, "Sí", "Por confirmar",
        "Staff: 3 noches por fase (dato de Jefatura de Cooperación). Días en sitio 4: confirmar. 12 beneficiarios + 12 acompañantes: confirmar."),
    tpl("Taller ENTRÉNATE", "Taller de 3 días con cuidadores y líderes en cabecera municipal.",
        2, 0, 2, 4, 3, 18, 0, 3, 2, 1, 1, 1, 2, "Completa", "TTE-PART-LOCAL", "KIT-ENT", "Sí", "Por confirmar",
        "EJEMPLO: 2 facilitadores que viajan (no estaba en PRESUP_1). Participantes varían por lugar: ajustar en Eventos."),
    tpl("Taller de arranque (equipo)", "Inducción del equipo implementador en Bogotá.",
        13, 3, 13, 5, 4, 0, 0, 0, 0, 0, 1, 0, 4, "Ninguna", None, None, "Sí", "Estimado",
        "13 viajan (4 coordinadores + 8 facilitadoras + jefe ERAE); 3 locales."),
    tpl("Curso ERAE Básico", "Curso de 4 días para líderes comunitarios.",
        0, 2, 0, 4, 0, 10, 0, 4, 4, 1, 1, 1, 4, "Completa", "TTE-PART-MUNIC", "KIT-PED", "Sí", "Estimado", "2 facilitadores locales."),
    tpl("Visita de seguimiento", "1 profesional viaja a territorio (monitoreo, género, coordinación).",
        1, 0, 1, 3, 2, 0, 0, 0, 0, 0, 1, 0, 0, "Ninguna", None, None, "Sí", "Por confirmar", "EJEMPLO: ajustar días/noches según el tipo de visita."),
    tpl("Taller comunitario ERAE / MCMT (paquete)", "Taller de 1 día contratado TODO INCLUIDO (salón, refrigerio, materiales). Facilitadoras pagadas en Personal.",
        0, 0, 0, 0, 0, 10, 0, 1, 0, 0, 1, 0, 0, "Ninguna", None, None, "No", "Por confirmar",
        "Una fila por departamento o municipio; Repeticiones = nº de talleres allí. Beneficiarios = 10 × talleres (personas distintas en cada taller).", pkg="TALLER-ERAE"),
]
krows = PROJ.get("plantillas", krows)
TK = Tbl(wsK, PK, kcols, krows)
dv_list(wsK, "=Lista_Alimentacion", TK.rng("Alimentación participantes"))
dv_list(wsK, "=DV_Tarifas", TK.rng("Tarifa transporte participantes")); dv_list(wsK, "=DV_Tarifas", TK.rng("Kit por participante"))
dv_list(wsK, "=DV_Tarifas", TK.rng("Paquete por evento"))
dv_list(wsK, "=Lista_SiNo", TK.rng("Mismos participantes en repeticiones"))
dv_list(wsK, "=Lista_Estados", TK.rng("Estado")); estado_cf(wsK, TK.rng("Estado"))
wsK.freeze_panes = "B5"

# ================================================================= PERFILES DE BENEFICIARIOS (new)
wsF = wb.create_sheet("Perfiles")
title(wsF, "Perfiles de beneficiarios – cómo se reparten por sexo, edad y discapacidad",
      "Cada actividad de la WBS elige un perfil. Mujeres + Hombres + Niñas + Niños debe sumar 100%. '% con discapacidad' es el % DENTRO de cada grupo. Si una actividad es distinta, copie el perfil con otro nombre y ajústelo.")
G4 = ["Mujeres", "Hombres", "Niñas", "Niños"]
PF = "tblPerfiles"
fcols = [dict(n="Perfil", w=34)]
for g in G4:
    fcols.append(dict(n=f"% {g}", w=10, fmt=PCT))
fcols.append(dict(n="Suma", w=8, fmt=PCT, f="+".join(R(PF, f"% {g}") for g in G4)))
for g in G4:
    fcols.append(dict(n=f"% {g} con discapacidad", w=13, fmt=PCT, note=f"Porcentaje de este grupo ({g.lower()}) que tiene discapacidad."))
fcols += [dict(n="Fuente", w=30), dict(n="Estado", w=13), dict(n="Notas", w=50, wrap=True)]
def PFL(n, m, h, na, no, dm=None, dh=None, dna=None, dno=None, fuente="Experiencia CIREC (Jefatura de Cooperación)", est="Estimado", notas=None):
    return {"Perfil": n, "% Mujeres": m, "% Hombres": h, "% Niñas": na, "% Niños": no,
            "% Mujeres con discapacidad": dm, "% Hombres con discapacidad": dh, "% Niñas con discapacidad": dna, "% Niños con discapacidad": dno,
            "Fuente": fuente, "Estado": est, "Notas": notas}
frows = PROJ.get("perfiles", [
    PFL("Brigadas de rehabilitación", .30, .60, .05, .05, 1, 1, 1, 1, notas="Beneficiarios de brigada = personas con discapacidad por definición (100%). Ajustar si se atiende a otros."),
    PFL("ENTRÉNATE", .70, .30, 0, 0, notas="No participan niñas ni niños. % con discapacidad: por estimar (cuidadores)."),
    PFL("Mi Cuerpo Mi Territorio (MCMT)", .60, .20, .10, .10, notas="% con discapacidad: por estimar."),
])
TF = Tbl(wsF, PF, fcols, frows)
dv_list(wsF, "=Lista_Estados", TF.rng("Estado")); estado_cf(wsF, TF.rng("Estado"))
wsF.conditional_formatting.add(TF.rng("Suma"), FormulaRule(formula=[f'AND(${TF.col("Perfil")}5<>"",ROUND(${TF.col("Suma")}5,4)<>1)'], fill=RED))
wsF.conditional_formatting.add(TF.rng("Suma"), FormulaRule(formula=[f'AND(${TF.col("Perfil")}5<>"",ROUND(${TF.col("Suma")}5,4)=1)'], fill=GRN))
name("DV_Perfiles", "tblPerfiles[Perfil]")

# ================================================================= MAPEO DONANTE
wsM = wb.create_sheet("MapeoDonante")
title(wsM, "Mapeo: bloque de costo interno → categoría de cada donante",
      "Para otro donante: renombre un encabezado (ej. GIZ → COSUDE) y llene su columna y su lista de categorías. Máximo 4 donantes a la vez.")
mcols = [dict(n="Bloque", w=30), dict(n="CBPF", w=42), dict(n="AECID", w=30), dict(n="UE", w=30), dict(n="GIZ", w=30), dict(n="Nota", w=60)]
mrows = [dict(r) for r in rows_of("tblMapeo")]
for r_m in mrows:
    if r_m["Bloque"] in PROJ.get("mapeo", {}):
        r_m["CBPF"], r_m["Nota"] = PROJ["mapeo"][r_m["Bloque"]]
if not any(r_m["Bloque"] == "Paquetes por evento" for r_m in mrows):
    mrows.append({"Bloque": "Paquetes por evento", "CBPF": "D. Contractual Services",
                  "Nota": "Servicio contratado todo incluido (salón, refrigerio, materiales). Confirmar la categoría con cada donante."})
TM = Tbl(wsM, "tblMapeo", mcols, mrows)
k0 = TM.r2 + 3
wsM.cell(row=k0 - 1, column=1, value="Categorías de cada donante, en el orden de su formato").font = H2
kc = [dict(n="Nº", w=6), dict(n="Categoría CBPF", w=42), dict(n="Categoría AECID", w=30), dict(n="Categoría UE", w=30), dict(n="Categoría GIZ", w=30)]
Tbl(wsM, "tblCategorias", kc, rows_of("tblCategorias"), hrow=k0)
wsM.freeze_panes = "B5"
BLOQUES = [r["Bloque"] for r in mrows]


def map_choose(key_ref):
    ch = "CHOOSE(DonIdx,tblMapeo[CBPF],tblMapeo[AECID],tblMapeo[UE],tblMapeo[GIZ])"
    return f'IFERROR(INDEX({ch},MATCH({key_ref},tblMapeo[Bloque],0)),"")'


# ================================================================= WBS
wsW = wb.create_sheet("WBS")
title(wsW, "WBS – estructura del proyecto (mismos códigos del marco lógico)",
      "Amarillo: lo que DICE la propuesta + la plantilla de evento de cada actividad. Gris: lo que quedó presupuestado. Si no coinciden → REVISAR.")
W = "tblWBS"
cod = R(W, "Código WBS"); wc = f'{cod}&".*"'
def two(fn):
    return f"({fn(cod)}+{fn(wc)})"
act = f'{R(W,"Nivel")}="Actividad"'
pctf = lambda col: f'IFERROR(INDEX(tblPerfiles[{col}],MATCH({R(W,"Perfil de beneficiarios")},tblPerfiles[Perfil],0)),0)'
wcols = [
    dict(n="Código WBS", w=10, fmt=TXT, note="Formato texto: 1.2.1, 1.10, etc. Use los códigos del marco lógico."),
    dict(n="Nivel", w=11), dict(n="Nombre", w=48), dict(n="Responsable técnico", w=18),
    dict(n="Plantilla", w=26, note="Plantilla de evento por defecto de esta actividad. Los eventos de la actividad la heredan."),
    dict(n="Eventos declarados", w=11, fmt="0", note="Cuántos eventos dice la propuesta (ej. 4 brigadas × 3 fases = 12)."),
    dict(n="Participantes declarados", w=12, fmt="0", note="Participantes ÚNICOS según la propuesta (no asistencias)."),
    dict(n="Perfil de beneficiarios", w=24, note="Cómo se reparten sus beneficiarios por sexo, edad y discapacidad (hoja Perfiles)."),
    dict(n="Beneficiarios (si difiere)", w=12, fmt=NUM, kind="opt", note="Vacío = participantes únicos de los eventos. Úselo cuando los beneficiarios no son eventos (ej. 7.600 asistentes a talleres comunitarios, 20 transferencias)."),
    dict(n="Suma al total", w=8, kind="opt", note="Vacío o Sí = cuenta en el total de beneficiarios del proyecto. No = no sumar (ej. las mismas personas ya contadas en otra actividad)."),
    dict(n="Código resultado", w=9, fmt=TXT, f=f'IF({cod}="","",LEFT({cod},FIND(".",{cod}&".")-1))'),
    dict(n="Código producto", w=9, fmt=TXT, f=f'IF(OR({cod}="",{R(W,"Nivel")}="Resultado"),"",LEFT({cod},IFERROR(FIND(".",{cod},FIND(".",{cod})+1)-1,LEN({cod}))))'),
    dict(n="Resultado", w=26, f=f'IFERROR(INDEX(tblWBS[Nombre],MATCH({R(W,"Código resultado")},tblWBS[Código WBS],0)),"")'),
    dict(n="Producto", w=26, f=f'IFERROR(INDEX(tblWBS[Nombre],MATCH({R(W,"Código producto")},tblWBS[Código WBS],0)),"")'),
    dict(n="Eventos presupuestados", w=11, fmt="0", f=f'IF({cod}="","",' + two(lambda k: f'SUMIFS(tblEventos[Eventos contados],tblEventos[Código WBS],{k})') + ')'),
    dict(n="Participantes presupuestados", w=12, fmt="0", f=f'IF({cod}="","",' + two(lambda k: f'SUMIFS(tblEventos[Participantes únicos],tblEventos[Código WBS],{k})') + ')'),
    dict(n="Costo eventos", w=14, fmt=PRES, f=f'IF({cod}="","",' + two(lambda k: f'SUMIFS(tblEventos[Total COP],tblEventos[Código WBS],{k})') + '*FX)'),
    dict(n="Costo personal", w=14, fmt=PRES, f=f'IF({cod}="","",' + two(lambda k: f'SUMIFS(tblPersonal[Total COP],tblPersonal[Código WBS],{k})') + '*FX)'),
    dict(n="Otros costos", w=14, fmt=PRES, f=f'IF({cod}="","",' + two(lambda k: f'SUMIFS(tblOtros[Total COP],tblOtros[Código WBS],{k})') + '*FX)'),
    dict(n="Total", w=15, fmt=PRES, f=f'IF({cod}="","",{R(W,"Costo eventos")}+{R(W,"Costo personal")}+{R(W,"Otros costos")})'),
    dict(n="Control", w=11, f=(f'IF(OR({cod}="",{R(W,"Nivel")}<>"Actividad"),"",IF(AND({R(W,"Eventos declarados")}="",{R(W,"Participantes declarados")}=""),'
                               f'IF({R(W,"Eventos presupuestados")}>0,"SIN METAS","—"),IF(OR(AND({R(W,"Eventos declarados")}<>"",{R(W,"Eventos declarados")}<>{R(W,"Eventos presupuestados")}),'
                               f'AND({R(W,"Participantes declarados")}<>"",{R(W,"Participantes declarados")}<>{R(W,"Participantes presupuestados")})),"REVISAR","OK")))')),
    dict(n="Notas", w=44),
    dict(n="Beneficiarios", w=11, fmt=NUM, f=f'IF(OR({cod}="",NOT({act})),"",IF({R(W,"Beneficiarios (si difiere)")}<>"",{R(W,"Beneficiarios (si difiere)")},{R(W,"Participantes presupuestados")}))'),
]
for g in G4:
    wcols.append(dict(n=f"p{g}", w=7, fmt=PCT, f=f'IF({R(W,"Perfil de beneficiarios")}="",0,N({pctf("% " + g)}))'))
for g in G4:
    wcols.append(dict(n=f"pd{g}", w=7, fmt=PCT, f=f'IF({R(W,"Perfil de beneficiarios")}="",0,N({pctf("% " + g + " con discapacidad")}))'))
_B = R(W, "Beneficiarios"); _p = {g: R(W, "p" + g) for g in G4}
_mx = "MAX(" + ",".join(_p.values()) + ")"
_r = {g: f"ROUND({_B}*{_p[g]},0)" for g in G4}
_adj = f"({_B}-(" + "+".join(_r.values()) + "))"
_first = {"Mujeres": f"{_p['Mujeres']}={_mx}",
          "Hombres": f"AND({_p['Hombres']}={_mx},{_p['Mujeres']}<{_mx})",
          "Niñas": f"AND({_p['Niñas']}={_mx},{_p['Mujeres']}<{_mx},{_p['Hombres']}<{_mx})",
          "Niños": f"AND({_p['Niños']}={_mx},{_p['Mujeres']}<{_mx},{_p['Hombres']}<{_mx},{_p['Niñas']}<{_mx})"}
for g in G4:
    wcols.append(dict(n=g, w=9, fmt=NUM, f=f'IF(OR({_B}="",{R(W,"Perfil de beneficiarios")}=""),"",{_r[g]}+IF({_first[g]},{_adj},0))'))
for g in G4:
    wcols.append(dict(n=f"{g} con discapacidad", w=11, fmt=NUM, f=f'IF({R(W, g)}="","",ROUND({R(W, g)}*{R(W, "pd" + g)},0))'))
wcols.append(dict(n="Total con discapacidad", w=11, fmt=NUM, f=f'IF({R(W,"Mujeres")}="","",' + "+".join(R(W, f"{g} con discapacidad") for g in G4) + ")"))
PL_BY_COD = {"1.2.1": "Brigada de rehabilitación (fase)", "1.2.2": "Taller ENTRÉNATE", "1.3.1": "Taller de arranque (equipo)",
             "1.3.2": "Curso ERAE Básico", "0.2": "Visita de seguimiento"}
wrows = []
for w in rows_of("tblWBS"):
    d = dict(w); d["Plantilla"] = PL_BY_COD.get(w["Código WBS"])
    d["Perfil de beneficiarios"] = {"1.2.1": "Brigadas de rehabilitación", "1.2.2": "ENTRÉNATE"}.get(w["Código WBS"]); wrows.append(d)
wrows.append({"Código WBS": "1.3.3", "Nivel": "Actividad", "Nombre": "Talleres comunitarios ERAE (750 talleres × 10 personas) – EJEMPLO",
              "Plantilla": "Taller comunitario ERAE / MCMT (paquete)", "Eventos declarados": 750, "Participantes declarados": 7500,
              "Perfil de beneficiarios": "Mi Cuerpo Mi Territorio (MCMT)", "Notas": "EJEMPLO de actividad masiva: 4 filas en Eventos (una por departamento) con Repeticiones = nº de talleres."})
wrows = PROJ.get("wbs", wrows)
TW = Tbl(wsW, W, wcols, wrows)
dv_list(wsW, "=Lista_Niveles", TW.rng("Nivel")); dv_list(wsW, "=DV_Plantillas", TW.rng("Plantilla"))
dv_list(wsW, "=DV_Perfiles", TW.rng("Perfil de beneficiarios")); dv_list(wsW, "=Lista_SiNo", TW.rng("Suma al total"))
wsW.column_dimensions.group(TW.col("pMujeres"), TW.col("pdNiños"), hidden=True, outline_level=1)
hcell(wsW, 3, TW.idx["Beneficiarios"], "BENEFICIARIOS (estimados con el perfil) →", HDRC)
ctl = TW.rng("Control")
wsW.conditional_formatting.add(ctl, CellIsRule(operator="equal", formula=['"REVISAR"'], fill=RED))
wsW.conditional_formatting.add(ctl, CellIsRule(operator="equal", formula=['"SIN METAS"'], fill=YEL))
wsW.conditional_formatting.add(ctl, CellIsRule(operator="equal", formula=['"OK"'], fill=GRN))
wsW.conditional_formatting.add(f"A5:C{4+GROW}", FormulaRule(formula=['$B5="Resultado"'], font=Font(bold=True, color="1F4E5A")))
money_cf(wsW, f"{TW.col('Costo eventos')}5:{TW.col('Total')}{4+GROW}")
wsW.freeze_panes = "D5"
wsW.column_dimensions.group(TW.col("Código resultado"), TW.col("Producto"), hidden=True, outline_level=1)

# ================================================================= EVENTOS
wsE = wb.create_sheet("Eventos")
title(wsE, "Eventos – una fila por cada LUGAR donde ocurre una actividad",
      "Mínimo: Código WBS + Lugar. Todo lo demás viene de la plantilla de la actividad. Si este evento es distinto, abra el grupo [+] 'Ajustes' y cambie solo esa celda.")
E = "tblEventos"; e = lambda c: R(E, c)
tplv = lambda p: f'INDEX(tblPlantillas[{p}],MATCH({e("Plantilla usada")},tblPlantillas[Plantilla],0))'
ecols = [
    dict(n="ID evento", w=7, fmt=TXT, f='"E"&TEXT(ROW()-4,"000")', note="Automático."),
    dict(n="Código WBS", w=9, fmt=TXT, note="OBLIGATORIO. Actividad a la que pertenece el evento."),
    dict(n="Actividad", w=26, f=f'IF({e("Código WBS")}="","",IFERROR(INDEX(tblWBS[Nombre],MATCH({e("Código WBS")},tblWBS[Código WBS],0)),"¡CÓDIGO NO EXISTE!"))'),
    dict(n="Plantilla usada", w=22, f=(f'IF({e("Ajuste plantilla")}<>"",{e("Ajuste plantilla")},IF({e("Código WBS")}="","",'
                                        f'IFERROR(T(INDEX(tblWBS[Plantilla],MATCH({e("Código WBS")},tblWBS[Código WBS],0))),"")))'),
         note="Sale de la columna Plantilla de la WBS. Para usar otra, cambie 'Ajuste plantilla'."),
    dict(n="Descripción", w=26, kind="opt", note="Opcional."),
    dict(n="Lugar", w=24, note="OBLIGATORIO. Elija de la hoja Lugares (si no está, agréguelo allá)."),
    dict(n="Departamento", w=13, f=f'IF({e("Lugar")}="","",IFERROR(INDEX(tblLugares[Departamento],MATCH({e("Lugar")},tblLugares[Lugar],0)),""))'),
    dict(n="Zona", w=13, f=f'IF({e("Ajuste zona")}<>"",{e("Ajuste zona")},IF({e("Lugar")}="","",IFERROR(INDEX(tblLugares[Zona],MATCH({e("Lugar")},tblLugares[Lugar],0)),"")))'),
    dict(n="Repeticiones", w=9, kind="opt", note="Vacío = 1. Veces que el MISMO evento se repite en el MISMO lugar (ej. 3 fases de brigada). Nunca use esto para lugares distintos: un lugar = una fila."),
    dict(n="Participantes (si difiere)", w=11, kind="opt", note="Vacío = el número de la plantilla."),
    dict(n="Beneficiarios", w=10, fmt=NUM, f='tblEventos[[#This Row],[Participantes únicos]]', note="Personas únicas que atiende este evento (participantes; sin acompañantes). Se cuentan una vez si las repeticiones son con las mismas personas."),
    dict(n="Mes", w=6, fmt="0", kind="opt", note="Opcional: mes del proyecto en que empieza (1, 2, 3...)."),
    dict(n="Año", w=6, fmt="0", f=f'IF({e("Mes")}="","",ROUNDUP({e("Mes")}/12,0))'),
    dict(n="Complementaria", w=9, kind="opt", note="Marque Sí solo si esta fila es un grupo extra de un evento que ya tiene su fila (mismo código y lugar). No cuenta como evento ni suma participantes."),
    dict(n="Estado", w=12, kind="opt"), dict(n="Notas", w=36, kind="opt", wrap=True),
]
ADJ = ["Ajuste plantilla", "Ajuste zona"] + [f"Ajuste {p.lower()}" for p in PARAMS_NUM if p != "Participantes"] + [f"Ajuste {p.lower()}" for p in PARAMS_TXT]
ADJ_OF = {p: f"Ajuste {p.lower()}" for p in PARAMS_NUM + PARAMS_TXT}
for a in ADJ:
    ecols.append(dict(n=a, w=10, kind="adj", fmt=PCT if "porcentaje" in a else None,
                      note="Vacío = valor de la plantilla. Escriba aquí solo si este evento es distinto."))
# effective values
for p in PARAMS_NUM:
    if p == "Participantes":
        ecols.append(dict(n="Participantes usados", w=9, f=f'IF({e("Participantes (si difiere)")}<>"",{e("Participantes (si difiere)")},IFERROR(N({tplv("Participantes")}),0))'))
    else:
        ecols.append(dict(n=p, w=9, fmt=PCT if p.startswith("Porcentaje") else None,
                          f=f'IF({e(ADJ_OF[p])}<>"",{e(ADJ_OF[p])},IFERROR(N({tplv(p)}),0))'))
for p in PARAMS_TXT:
    ecols.append(dict(n=p, w=11, f=f'IF({e(ADJ_OF[p])}<>"",{e(ADJ_OF[p])},IFERROR(T({tplv(p)}),""))'))
RU = "Repeticiones usadas"
ecols += [
    dict(n=RU, w=9, f=f'IF({e("Repeticiones")}="",1,{e("Repeticiones")})'),
    dict(n="Eventos contados", w=9, f=f'IF(OR({e("Código WBS")}="",{e("Complementaria")}="Sí"),0,{e(RU)})'),
    dict(n="Participantes únicos", w=9, f=f'IF({e("Complementaria")}="Sí",0,{e("Participantes usados")}*IF({e("Mismos participantes en repeticiones")}="No",{e(RU)},1))'),
]
P_ = e("Participantes usados"); A_ = e("Acompañantes"); r_ = e(RU)
blocks = [
    ("Hospedaje staff", "noche-persona", f'{e("Staff viaja")}*{e("Noches staff")}*{r_}', rate('"HOSP-STAFF"', e("Zona"))),
    ("Alimentación staff", "persona-día", f'({e("Staff viaja")}+{e("Staff local")})*{e("Días staff en sitio")}*{r_}', rate('"ALIM"', e("Zona"))),
    ("Tiquetes aéreos staff", "tiquete", f'{e("Staff en avión")}*{r_}', rate('"TIQ"', e("Zona"))),
    ("Transporte terrestre staff", "persona-día", f'{e("Staff viaja")}*{e("Días transporte staff")}*{r_}', rate('"TTE-STAFF"', e("Zona"))),
    ("Hospedaje participantes", "habitación-noche",
     f'IF({P_}+{A_}=0,0,ROUNDUP(({P_}+{A_})*{e("Porcentaje con hospedaje")}/MAX({e("Personas por habitación")},1),0)*{e("Noches participantes")}*{r_})',
     rate('"HOSP-PART"', e("Zona"))),
    ("Alimentación participantes", "persona-día",
     f'IF({e("Alimentación participantes")}="Completa",({P_}+{A_})*{e("Días participantes")}*{r_},IF({e("Alimentación participantes")}="Refrigerio",{P_}*{e("Días participantes")}*{r_},0))',
     f'IF({e("Alimentación participantes")}="Refrigerio",{rate(chr(34)+"REFRIG"+chr(34), e("Zona"))},{rate(chr(34)+"ALIM"+chr(34), e("Zona"))})'),
    ("Transporte participantes", "trayecto", f'ROUND(({P_}+{A_})*{e("Porcentaje con transporte")},0)*{r_}', rate(e("Tarifa transporte participantes"), e("Zona"))),
    ("Alquiler de salones", "día", f'{e("Días de salón")}*{r_}', rate('"SALON"', e("Zona"))),
    ("Materiales y kits", "kit", f'IF({e("Kit por participante")}="",0,{P_}*IF({e("Mismos participantes en repeticiones")}="No",{r_},1))',
     rate(e("Kit por participante"), e("Zona"))),
    ("Paquetes por evento", "paquete", f'IF(OR({e("Paquete por evento")}="",{e("Complementaria")}="Sí"),0,{r_})', rate(e("Paquete por evento"), e("Zona"))),
]
for b in blocks:
    ecols.append(dict(n=f"Q {b[0]}", w=10, fmt=NUM, kind="qty", f=b[2], note=f"Cantidad en {b[1]}"))
for b in blocks:
    ecols.append(dict(n=f"COP {b[0]}", w=12, fmt=COP, f=f'IF({e("Q " + b[0])}=0,0,{e("Q " + b[0])}*{b[3]})'))
ecols.append(dict(n="Total COP", w=14, fmt=COP, f="SUM(" + ",".join(e(f"COP {b[0]}") for b in blocks) + ")"))
ecols.append(dict(n="Total moneda", w=13, fmt=PRES, f=f'{e("Total COP")}*FX', note="En la moneda de presentación elegida en Parametros."))

evrows = []
for lug in ["San José del Guaviare", "Florencia", "Popayán", "Medellín"]:
    evrows.append({"Código WBS": "1.2.1", "Lugar": lug, "Repeticiones": 3, "Estado": "Por confirmar",
                   "Notas": "EJEMPLO. 3 fases = 3 repeticiones en el mismo lugar." if lug.startswith("San") else None})
for dep, p in [("Guaviare", 12), ("Caquetá", 15), ("Cauca", 22), ("Antioquia", 23)]:
    evrows.append({"Código WBS": "1.2.2", "Lugar": f"Por definir ({dep})", "Participantes (si difiere)": p, "Estado": "Por confirmar",
                   "Notas": "EJEMPLO. Participantes distintos por lugar → columna 'Participantes (si difiere)'." if dep == "Guaviare" else None})
evrows.append({"Código WBS": "1.3.1", "Lugar": "Bogotá", "Mes": 1, "Estado": "Estimado"})
evrows.append({"Código WBS": "1.3.2", "Lugar": "Popayán", "Estado": "Estimado"})
for dep, n in [("Guaviare", 200), ("Caquetá", 190), ("Cauca", 180), ("Antioquia", 180)]:
    evrows.append({"Código WBS": "1.3.3", "Lugar": f"Por definir ({dep})", "Descripción": f"Talleres comunitarios ERAE {dep}", "Repeticiones": n,
                   "Estado": "Por confirmar", "Notas": "EJEMPLO: 1 fila por departamento; Repeticiones = talleres en el departamento." if dep == "Guaviare" else None})
evrows = PROJ.get("eventos", evrows)
TE = Tbl(wsE, E, ecols, evrows)
dv_list(wsE, "=DV_WBS", TE.rng("Código WBS")); dv_list(wsE, "=DV_Lugares", TE.rng("Lugar"))
dv_list(wsE, "=Lista_Si", TE.rng("Complementaria")); dv_list(wsE, "=Lista_Estados", TE.rng("Estado"))
dv_list(wsE, "=DV_Plantillas", TE.rng("Ajuste plantilla")); dv_list(wsE, "=Lista_Zonas", TE.rng("Ajuste zona"))
dv_list(wsE, "=Lista_Alimentacion", TE.rng(ADJ_OF["Alimentación participantes"]))
dv_list(wsE, "=DV_Tarifas", TE.rng(ADJ_OF["Tarifa transporte participantes"])); dv_list(wsE, "=DV_Tarifas", TE.rng(ADJ_OF["Kit por participante"]))
dv_list(wsE, "=DV_Tarifas", TE.rng(ADJ_OF["Paquete por evento"]))
dv_list(wsE, "=Lista_SiNo", TE.rng(ADJ_OF["Mismos participantes en repeticiones"]))
estado_cf(wsE, TE.rng("Estado"))
wsE.conditional_formatting.add(TE.rng("Actividad"), CellIsRule(operator="equal", formula=['"¡CÓDIGO NO EXISTE!"'], fill=RED))
wsE.conditional_formatting.add(TE.rng("Departamento"), FormulaRule(formula=[f'AND(${TE.col("Lugar")}5<>"",${TE.col("Departamento")}5="")'], fill=RED))
wsE.conditional_formatting.add(TE.rng("Plantilla usada"), FormulaRule(formula=[f'AND(${TE.col("Código WBS")}5<>"",${TE.col("Plantilla usada")}5="")'], fill=RED))
money_cf(wsE, TE.rng("Total moneda"))
# highlight adjusted cells (any override filled) in the visible area? -> mark row with a flag column-free CF on Notas
first_adj, last_adj = TE.col(ADJ[0]), TE.col(ADJ[-1])
first_calc, last_calc = TE.col("Participantes usados"), TE.col(f"COP {blocks[-1][0]}")
wsE.column_dimensions.group(first_adj, last_adj, hidden=True, outline_level=1)
wsE.column_dimensions.group(first_calc, last_calc, hidden=True, outline_level=2)
wsE.conditional_formatting.add(f"A5:A{4+GROW}", FormulaRule(formula=[f'COUNTA(${first_adj}5:${last_adj}5)>0'], fill=PatternFill("solid", fgColor="F4D9A6")))
wsE.freeze_panes = "D5"
hcell(wsE, 3, TE.idx[ADJ[0]], "AJUSTES (solo si el evento difiere de la plantilla)", HDRA)
hcell(wsE, 3, TE.idx["Participantes usados"], "VALORES USADOS (calculados)", HDRC)
hcell(wsE, 3, TE.idx["Q " + blocks[0][0]], "CANTIDADES", HDRC)
hcell(wsE, 3, TE.idx["COP " + blocks[0][0]], "COSTOS COP", HDRC)

# ================================================================= PERSONAL
wsR = wb.create_sheet("Personal")
title(wsR, "Personal – cargo × número × meses × dedicación", "El costo mensual viene de Cargos. 'Cálculo' se escribe solo: cópielo a la narrativa del donante.")
Pn = "tblPersonal"; p = lambda c: R(Pn, c)
pcols = [
    dict(n="ID", w=7, fmt=TXT, f='"P"&TEXT(ROW()-4,"000")'), dict(n="Código WBS", w=9, fmt=TXT), dict(n="Cargo", w=36),
    dict(n="Número de personas", w=9), dict(n="Meses", w=7), dict(n="Dedicación", w=10, fmt=PCT, note="Como porcentaje: 15%."),
    dict(n="Costo mensual COP", w=15, fmt=COP, f=f'IF({p("Cargo")}="","",IFERROR(INDEX(tblCargos[Costo mensual COP],MATCH({p("Cargo")},tblCargos[Cargo],0)),0))'),
    dict(n="Total COP", w=15, fmt=COP, f=f'IF({p("Cargo")}="","",{p("Número de personas")}*{p("Meses")}*{p("Dedicación")}*{p("Costo mensual COP")})'),
    dict(n="Total moneda", w=13, fmt=PRES, f=f'IF({p("Cargo")}="","",{p("Total COP")}*FX)'),
    dict(n="Cálculo", w=70, wrap=True, f=(f'IF({p("Cargo")}="","",{p("Número de personas")}&" ("&{p("Cargo")}&") × "&{mtxt(p("Costo mensual COP") + "*FX")}&"/mes × "'
                                          f'&{p("Meses")}&" meses × "&FIXED({p("Dedicación")}*100,0)&"% dedicación = "&{mtxt(p("Total moneda"))})')),
    dict(n="Rol y justificación", w=44, wrap=True, kind="opt"), dict(n="Estado", w=12, kind="opt"),
]
TP = Tbl(wsR, Pn, pcols, PROJ.get("personal", rows_of("tblPersonal")))
dv_list(wsR, "=DV_WBS", TP.rng("Código WBS")); dv_list(wsR, "=DV_Cargos", TP.rng("Cargo")); dv_list(wsR, "=Lista_Estados", TP.rng("Estado"))
estado_cf(wsR, TP.rng("Estado")); money_cf(wsR, TP.rng("Total moneda"))
wsR.freeze_panes = "D5"

# ================================================================= OTROS COSTOS
wsO = wb.create_sheet("OtrosCostos")
title(wsO, "Otros costos – lo que no depende de un evento (dispositivos, kits masivos, seguros, equipos, operación)",
      "Cantidad × Factor × Tarifa. El Bloque sale solo de la tarifa; use 'Bloque (si difiere)' solo para excepciones.")
On = "tblOtros"; o = lambda c: R(On, c)
zo = f'IF({o("Zona")}="","Todas",{o("Zona")})'
ocols = [
    dict(n="ID", w=7, fmt=TXT, f='"O"&TEXT(ROW()-4,"000")'), dict(n="Código WBS", w=9, fmt=TXT), dict(n="ID tarifa", w=14, fmt=TXT),
    dict(n="Concepto", w=36, f=f'IF({o("ID tarifa")}="","",IFERROR(INDEX(tblTarifas[Concepto],MATCH({o("ID tarifa")},tblTarifas[ID tarifa],0)),"¡TARIFA NO EXISTE!"))'),
    dict(n="Zona", w=13, kind="opt", note="Vacío = Todas."), dict(n="Cantidad", w=9, fmt="#,##0"), dict(n="Factor", w=8, note="Meses, fases, noches... Vacío no: escriba 1 si no aplica."),
    dict(n="Unidad del factor", w=11, kind="opt"),
    dict(n="Tarifa COP", w=13, fmt=COP, f=f'IF({o("ID tarifa")}="","",{rate(o("ID tarifa"), zo)})'),
    dict(n="Total COP", w=15, fmt=COP, f=f'IF({o("ID tarifa")}="","",{o("Cantidad")}*{o("Factor")}*{o("Tarifa COP")})'),
    dict(n="Total moneda", w=13, fmt=PRES, f=f'IF({o("ID tarifa")}="","",{o("Total COP")}*FX)'),
    dict(n="Bloque (si difiere)", w=18, kind="opt", note="Vacío = el bloque de la tarifa (hoja Tarifas)."),
    dict(n="Bloque", w=24, f=f'IF({o("Bloque (si difiere)")}<>"",{o("Bloque (si difiere)")},IF({o("ID tarifa")}="","",IFERROR(T(INDEX(tblTarifas[Bloque],MATCH({o("ID tarifa")},tblTarifas[ID tarifa],0))),"")))'),
    dict(n="Cálculo", w=70, wrap=True, f=(f'IF({o("ID tarifa")}="","",{o("Cantidad")}&" ("&{o("Concepto")}&") × "&{o("Factor")}&" "&{o("Unidad del factor")}&" × "'
                                          f'&{mtxt(o("Tarifa COP") + "*FX")}&" = "&{mtxt(o("Total moneda"))})')),
    dict(n="Justificación", w=40, wrap=True, kind="opt"), dict(n="Estado", w=12, kind="opt"),
]
orows = []
for x in rows_of("tblOtros"):
    d = dict(x); d["Zona"] = None if x.get("Zona") == "Todas" else x.get("Zona")
    default = BLOQUE_POR_ID.get(x["ID tarifa"])
    d["Bloque (si difiere)"] = None if default == x.get("Bloque") else x.get("Bloque")
    if x["ID tarifa"] != "REFRIG":
        orows.append(d)
orows = PROJ.get("otros", orows)
TO = Tbl(wsO, On, ocols, orows)
dv_list(wsO, "=DV_WBS", TO.rng("Código WBS")); dv_list(wsO, "=DV_Tarifas", TO.rng("ID tarifa")); dv_list(wsO, "=Lista_Zonas", TO.rng("Zona"))
dv_list(wsO, "=DV_Bloques", TO.rng("Bloque (si difiere)")); dv_list(wsO, "=Lista_Estados", TO.rng("Estado"))
estado_cf(wsO, TO.rng("Estado")); money_cf(wsO, TO.rng("Total moneda"))
wsO.conditional_formatting.add(TO.rng("Concepto"), CellIsRule(operator="equal", formula=['"¡TARIFA NO EXISTE!"'], fill=RED))
wsO.freeze_panes = "D5"

# ================================================================= RESUMEN
wsS = wb.create_sheet("Resumen")
title(wsS, "Resumen del presupuesto", "Solo cálculos. Muestra siempre COP y USD.")
for i, (h, w) in enumerate([("Bloque de costo", 32), ("Categoría del donante activo", 44), ("Total COP", 18), ("Total USD", 15), ("% del directo", 11)], 1):
    hcell(wsS, 4, i, h); wsS.column_dimensions[L(i)].width = w
ev_blocks = {b[0] for b in blocks}
S1 = 5; S2 = S1 + len(BLOQUES) - 1
for i, bn in enumerate(BLOQUES):
    r = S1 + i
    bcell(wsS, r, 1, bn)
    bcell(wsS, r, 2, f'=IF({map_choose(f"A{r}")}="","SIN MAPEO",{map_choose(f"A{r}")})')
    parts = [f'SUMIF(tblOtros[Bloque],A{r},tblOtros[Total COP])']
    if bn == "Personal": parts.append("SUM(tblPersonal[Total COP])")
    if bn in ev_blocks: parts.append(f"SUM(tblEventos[COP {bn}])")
    bcell(wsS, r, 3, "=" + "+".join(parts), COP)
    bcell(wsS, r, 4, f"=C{r}/TRM", USD)
    bcell(wsS, r, 5, f"=IF($C${S2+1}=0,0,C{r}/$C${S2+1})", "0.0%")
wsS.conditional_formatting.add(f"B{S1}:B{S2}", CellIsRule(operator="equal", formula=['"SIN MAPEO"'], fill=RED))
rt = S2 + 1
wsS.cell(row=rt, column=1, value="TOTAL COSTOS DIRECTOS").font = BOLD
bcell(wsS, rt, 3, f"=SUM(C{S1}:C{S2})", COP, font=BOLD); bcell(wsS, rt, 4, f"=C{rt}/TRM", USD, font=BOLD)
wsS.cell(row=rt + 1, column=1, value="Costos indirectos").font = BF
wsS.cell(row=rt + 1, column=2, value='="Tasa "&FIXED(PctInd*100,1)&"% sobre "&BaseInd').font = SUB
bcell(wsS, rt + 1, 3, f'=IF(BaseInd="Total del proyecto",C{rt}*PctInd/(1-PctInd),C{rt}*PctInd)', COP)
bcell(wsS, rt + 1, 4, f"=C{rt+1}/TRM", USD)
wsS.cell(row=rt + 2, column=1, value="TOTAL PROYECTO").font = BOLD
bcell(wsS, rt + 2, 3, f"=C{rt}+C{rt+1}", COP, font=BOLD); bcell(wsS, rt + 2, 4, f"=C{rt+2}/TRM", USD, font=BOLD)
name("Res_DirectoCOP", f"Resumen!$C${rt}"); name("Res_IndCOP", f"Resumen!$C${rt+1}"); name("Res_TotalCOP", f"Resumen!$C${rt+2}")
name("Res_TotalUSD", f"Resumen!$D${rt+2}")
wsS.cell(row=rt + 3, column=1, value="Margen frente al techo del donante").font = BF
wsS.cell(row=rt + 3, column=2, value='="en "&Moneda').font = SUB
bcell(wsS, rt + 3, 4, '=IF(Techo="","",Techo-Res_TotalCOP*FX)', COP); money_cf(wsS, f"D{rt+3}")
r0 = rt + 6
wsS.cell(row=r0 - 1, column=1, value="Por categoría del donante activo").font = H2
for i, h in enumerate(["Nº", "Categoría", "Total COP", "Total USD", "%"], 1):
    hcell(wsS, r0, i, h)
chc = "CHOOSE(DonIdx,tblCategorias[Categoría CBPF],tblCategorias[Categoría AECID],tblCategorias[Categoría UE],tblCategorias[Categoría GIZ])"
for i in range(10):
    r = r0 + 1 + i
    bcell(wsS, r, 1, i + 1)
    bcell(wsS, r, 2, f'=IFERROR(IF(INDEX({chc},{i+1})="","",INDEX({chc},{i+1})),"")')
    bcell(wsS, r, 3, f'=IF(B{r}="","",SUMIF($B${S1}:$B${S2},B{r},$C${S1}:$C${S2}))', COP)
    bcell(wsS, r, 4, f'=IF(B{r}="","",C{r}/TRM)', USD)
    bcell(wsS, r, 5, f'=IF(B{r}="","",IF(Res_DirectoCOP=0,0,C{r}/Res_DirectoCOP))', "0.0%")
rc = r0 + 11
wsS.cell(row=rc, column=2, value="Total por categorías (debe ser igual al directo)").font = BOLD
bcell(wsS, rc, 3, f"=SUM(C{r0+1}:C{r0+10})", COP, font=BOLD)
name("Res_CatCOP", f"Resumen!$C${rc}")
name("Res_CatNom", f"Resumen!$B${r0+1}:$B${r0+10}"); name("Res_CatVal", f"Resumen!$C${r0+1}:$C${r0+10}")
ry = rc + 3
wsS.cell(row=ry - 1, column=1, value="Costo de eventos por año del proyecto (según columna Mes)").font = H2
for i, h in enumerate(["Año", "Nº eventos", "Total COP", "Total USD"], 1):
    hcell(wsS, ry, i, h)
for i, y in enumerate([1, 2, 3, 4, "Sin mes"]):
    r = ry + 1 + i
    bcell(wsS, r, 1, y)
    crit = '""' if y == "Sin mes" else f"A{r}"
    bcell(wsS, r, 2, f'=SUMIFS(tblEventos[Eventos contados],tblEventos[Año],{crit})', NUM)
    bcell(wsS, r, 3, f"=SUMIFS(tblEventos[Total COP],tblEventos[Año],{crit})", COP)
    bcell(wsS, r, 4, f"=C{r}/TRM", USD)

# ================================================================= LINEAS DONANTE (compact: only activities)
wsD = wb.create_sheet("LineasDonante")
title(wsD, "Líneas para el formato del donante (eventos agregados por actividad × bloque)",
      "Filtre la columna Total ≠ 0. Personal y OtrosCostos ya tienen su propia columna 'Cálculo'.")
dh = ["Código WBS", "Actividad", "Bloque", "Categoría (donante activo)", "Cantidad", "Unidad", "Costo unitario promedio", "Total", "Lugares", "Cálculo"]
for i, (h, w) in enumerate(zip(dh, [10, 40, 26, 36, 11, 16, 14, 14, 9, 90]), 1):
    hcell(wsD, 4, i, h); wsD.column_dimensions[L(i)].width = w
ACT_K = lambda k: (f'IFERROR(INDEX(tblWBS[Código WBS],SMALL(IF(tblWBS[Nivel]="Actividad",'
                   f'ROW(tblWBS[Nivel])-MIN(ROW(tblWBS[Nivel]))+1),{k})),"")')
r = 5
for k in range(1, MAX_ACT_LD + 1):
    for j, b in enumerate(blocks):
        if j == 0:
            wsD[f"A{r}"] = ArrayFormula(f"A{r}", "=" + ACT_K(k)); first = r
        else:
            wsD[f"A{r}"] = f"=A{first}"
        wsD.cell(row=r, column=2, value=f'=IF(A{r}="","",INDEX(tblWBS[Nombre],MATCH(A{r},tblWBS[Código WBS],0)))')
        wsD.cell(row=r, column=3, value=f'=IF(A{r}="","","{b[0]}")')
        wsD.cell(row=r, column=4, value=f'=IF(A{r}="","",{map_choose(f"C{r}")})')
        wsD.cell(row=r, column=5, value=f'=IF(A{r}="","",SUMIFS(tblEventos[Q {b[0]}],tblEventos[Código WBS],A{r}))')
        wsD.cell(row=r, column=6, value=f'=IF(A{r}="","","{b[1]}")')
        wsD.cell(row=r, column=8, value=f'=IF(A{r}="","",SUMIFS(tblEventos[COP {b[0]}],tblEventos[Código WBS],A{r})*FX)')
        wsD.cell(row=r, column=7, value=f'=IF(A{r}="","",IF(E{r}=0,0,H{r}/E{r}))')
        wsD.cell(row=r, column=9, value=f'=IF(A{r}="","",COUNTIFS(tblEventos[Código WBS],A{r},tblEventos[Q {b[0]}],">0"))')
        wsD.cell(row=r, column=10, value=(f'=IF(OR(A{r}="",H{r}=0),"",FIXED(E{r},0)&" "&F{r}&" × "&{mtxt(f"G{r}")}&" (promedio) = "'
                                           f'&{mtxt(f"H{r}")}&". Desglose en "&I{r}&" fila(s) de la actividad "&A{r}&" (hoja Eventos).")'))
        for c in range(1, 11):
            x = wsD.cell(row=r, column=c); x.font = BF; x.border = BRD; x.fill = CAL
        wsD.cell(row=r, column=5).number_format = '#,##0'
        wsD.cell(row=r, column=7).number_format = PRES; wsD.cell(row=r, column=8).number_format = PRES
        r += 1
DL2 = r - 1
wsD.auto_filter.ref = f"A4:J{DL2}"
money_cf(wsD, f"G5:H{DL2}")
wsD.freeze_panes = "C5"

# ================================================================= AUDITORIA
wsA = wb.create_sheet("Auditoria")
title(wsA, "Auditoría automática – revisar ANTES de pasar al formato del donante",
      "ROJO = corregir. AMARILLO = confirmar con el equipo técnico. VERDE = pasó. Solo detecta lo que está programado: no reemplaza la revisión humana.")
for i, (h, w) in enumerate([("#", 5), ("Chequeo", 62), ("Resultado", 12), ("Estado", 12), ("Qué hacer", 72)], 1):
    hcell(wsA, 6, i, h); wsA.column_dimensions[L(i)].width = w


def notf(tbl, cols):
    return "+".join(f"SUMPRODUCT(--NOT(_xlfn.ISFORMULA({tbl}[{c}])))" for c in cols)


Ev = lambda c: f"tblEventos[{c}]"
ev_calc = [c["n"] for c in ecols if c.get("f")]
main = f'({Ev("Complementaria")}<>"Sí")'
checks = [
    ("WBS / METAS", None, None, None),
    ("Actividades cuyo nº de eventos o participantes presupuestados NO coincide con lo declarado", 'COUNTIF(tblWBS[Control],"REVISAR")', "ROJO", "Hoja WBS (o Tablero), filas en rojo: falta o sobra un evento, o los participantes no cuadran."),
    ("Actividades con eventos pero sin metas declaradas (no se pueden verificar)", 'COUNTIF(tblWBS[Control],"SIN METAS")', "AMARILLO", "Escribir en la WBS cuántos eventos/participantes dice la propuesta."),
    ("Códigos WBS repetidos", 'SUMPRODUCT((tblWBS[Código WBS]<>"")*(COUNTIF(tblWBS[Código WBS],tblWBS[Código WBS])>1))', "ROJO", "Cada código debe aparecer una sola vez."),
    (f"Más actividades de las que caben en Tablero ({MAX_ACT_TAB}) o LineasDonante ({MAX_ACT_LD})", f'MAX(0,COUNTIF(tblWBS[Nivel],"Actividad")-{MAX_ACT_TAB})', "ROJO", "Pida ampliar el tablero: las actividades sobrantes no se verían."),
    ("INTEGRIDAD", None, None, None),
    ("Celdas calculadas (grises) sobrescritas a mano",
     "+".join([notf("tblEventos", ev_calc), notf("tblPersonal", ["ID", "Costo mensual COP", "Total COP", "Total moneda", "Cálculo"]),
               notf("tblOtros", ["ID", "Concepto", "Tarifa COP", "Total COP", "Total moneda", "Bloque", "Cálculo"]),
               notf("tblWBS", ["Eventos presupuestados", "Participantes presupuestados", "Costo eventos", "Costo personal", "Otros costos", "Total", "Control", "Beneficiarios", "Mujeres", "Hombres", "Niñas", "Niños", "Total con discapacidad"]),
               notf("tblTarifas", ["Valor COP", "Valor USD", "En uso"])]),
     "ROJO", "Alguien escribió encima de una fórmula. Copie la fórmula de la celda de arriba (o Ctrl+Z)."),
    ("Total de Resumen vs. suma directa de Eventos + Personal + OtrosCostos (COP)",
     "ROUND(ABS(Res_DirectoCOP-(SUM(tblEventos[Total COP])+SUM(tblPersonal[Total COP])+SUM(tblOtros[Total COP]))),0)", "ROJO", "Hay un costo que no llega al Resumen (fila sin bloque o bloque mal escrito)."),
    ("Total en USD inconsistente con total en COP", "ROUND(ABS(Res_TotalUSD*TRM-Res_TotalCOP),0)", "ROJO", "Error interno de fórmulas en Resumen."),
    ("Bloques sin categoría en el donante activo", f'COUNTIF(Resumen!$B${S1}:$B${S2},"SIN MAPEO")', "ROJO", "Completar la columna del donante en MapeoDonante."),
    ("Diferencia entre total por bloques y total por categorías (COP)", "ROUND(ABS(Res_DirectoCOP-Res_CatCOP),0)", "ROJO", "Hay costos que no caen en ninguna categoría del donante."),
    ("EVENTOS", None, None, None),
    ("Eventos con código WBS inexistente", f'SUMPRODUCT(({Ev("Código WBS")}<>"")*(COUNTIF(tblWBS[Código WBS],{Ev("Código WBS")})=0))', "ROJO", "Corregir el código."),
    ("Eventos asignados a un código que no es de nivel Actividad",
     f'SUMPRODUCT(({Ev("Código WBS")}<>"")*(COUNTIFS(tblWBS[Código WBS],{Ev("Código WBS")},tblWBS[Nivel],"Actividad")=0)*(COUNTIF(tblWBS[Código WBS],{Ev("Código WBS")})>0))', "ROJO", "Presupuestar solo contra actividades (ej. 1.2.1)."),
    ("Eventos sin plantilla (la actividad no tiene plantilla en la WBS o el nombre no existe)",
     f'SUMPRODUCT(({Ev("Código WBS")}<>"")*(COUNTIF(tblPlantillas[Plantilla],{Ev("Plantilla usada")})=0))', "ROJO", "Asignar una plantilla a la actividad en la WBS (columna Plantilla)."),
    ("Eventos sin lugar o con lugar que no existe en Lugares",
     f'SUMPRODUCT(({Ev("Código WBS")}<>"")*(COUNTIF(tblLugares[Lugar],{Ev("Lugar")})=0))', "ROJO", "Elegir el lugar de la lista; si no existe, agregarlo en la hoja Lugares."),
    ("Repeticiones inválidas (menores a 1 o con decimales)", f'SUMPRODUCT(({Ev("Código WBS")}<>"")*((({Ev(RU)}<1)+({Ev(RU)}<>INT({Ev(RU)})))>0))', "ROJO", "Repeticiones = número entero ≥ 1."),
    ("Filas con cantidad pero sin tarifa (costo = 0)", "+".join(f'SUMPRODUCT(({Ev("Q " + b[0])}>0)*({Ev("COP " + b[0])}=0))' for b in blocks), "ROJO", "Falta la tarifa en Tarifas para ese ID / zona."),
    ("Staff que viaja más de 1 día con 0 noches de hospedaje", f'SUMPRODUCT(({Ev("Staff viaja")}>0)*({Ev("Días staff en sitio")}>1)*({Ev("Noches staff")}=0))', "ROJO", "Olvido típico de hospedaje. Si regresan cada día, poner 1 día."),
    ("Staff en avión mayor que staff que viaja", f'SUMPRODUCT(({Ev("Staff en avión")}>{Ev("Staff viaja")})*1)', "ROJO", "Solo vuela quien viaja."),
    ("Noches mayores que días (staff o participantes)", f'SUMPRODUCT(({Ev("Noches staff")}>{Ev("Días staff en sitio")})*1)+SUMPRODUCT(({Ev("Noches participantes")}>{Ev("Días participantes")})*1)', "AMARILLO", "Convención: noches normalmente = días − 1."),
    ("Eventos con participantes pero sin personal CIREC (¿quién facilita?)", f'SUMPRODUCT(({Ev("Participantes usados")}>0)*(({Ev("Staff viaja")}+{Ev("Staff local")})=0)*{main}*({Ev("Paquete por evento")}=""))', "AMARILLO", "Confirmar si facilita personal pagado en otra línea."),
    ("Participantes con días de asistencia pero sin alimentación", f'SUMPRODUCT(({Ev("Participantes usados")}>0)*({Ev("Días participantes")}>0)*((({Ev("Alimentación participantes")}="Ninguna")+({Ev("Alimentación participantes")}=""))>0))', "AMARILLO", "Confirmar si es intencional."),
    ("Filas 'Complementaria' sin fila principal (misma actividad y lugar)",
     f'SUMPRODUCT(({Ev("Complementaria")}="Sí")*(COUNTIFS({Ev("Código WBS")},{Ev("Código WBS")},{Ev("Lugar")},{Ev("Lugar")},{Ev("Complementaria")},"<>Sí")=0))', "ROJO", "La fila complementaria debe tener el mismo código y lugar que su evento principal."),
    ("Posibles eventos duplicados (misma actividad, lugar, descripción y mes)",
     f'SUMPRODUCT(({Ev("Código WBS")}<>"")*{main}*(COUNTIFS({Ev("Código WBS")},{Ev("Código WBS")},{Ev("Lugar")},{Ev("Lugar")},{Ev("Descripción")},{Ev("Descripción")}&"",{Ev("Mes")},{Ev("Mes")}&"",{Ev("Complementaria")},"<>Sí")>1))',
     "AMARILLO", "Si es el mismo evento repetido, use Repeticiones en una sola fila. Si son distintos, póngales descripción o mes."),
    ("Eventos sin estado o 'Por confirmar'", f'COUNTIFS({Ev("Código WBS")},"<>",{Ev("Estado")},"Por confirmar")+COUNTIFS({Ev("Código WBS")},"<>",{Ev("Estado")},"")', "AMARILLO", "Confirmar con el responsable técnico."),
    ("PLANTILLAS, LUGARES Y BENEFICIARIOS", None, None, None),
    ("Plantillas en uso no confirmadas", 'SUMPRODUCT((tblPlantillas[Plantilla]<>"")*(COUNTIF(tblEventos[Plantilla usada],tblPlantillas[Plantilla])>0)*(tblPlantillas[Estado]<>"Confirmado"))', "AMARILLO", "Validar cada plantilla con el equipo técnico: afecta todos sus eventos."),
    ("Plantillas con staff en avión mayor que staff que viaja", 'SUMPRODUCT((tblPlantillas[Staff en avión]>tblPlantillas[Staff viaja])*1)', "ROJO", "Corregir la plantilla."),
    ("Plantillas o lugares con nombre repetido", 'SUMPRODUCT((tblPlantillas[Plantilla]<>"")*(COUNTIF(tblPlantillas[Plantilla],tblPlantillas[Plantilla])>1))+SUMPRODUCT((tblLugares[Lugar]<>"")*(COUNTIF(tblLugares[Lugar],tblLugares[Lugar])>1))', "ROJO", "Cada nombre debe ser único (ej. 'Villanueva (Casanare)')."),
    ("Perfiles de beneficiarios que no suman 100% (mujeres + hombres + niñas + niños)",
     'SUMPRODUCT((tblPerfiles[Perfil]<>"")*(ROUND(tblPerfiles[Suma],4)<>1))', "ROJO", "Corregir los % en la hoja Perfiles."),
    ("Perfiles con % de discapacidad mayor a 100%",
     'COUNTIF(tblPerfiles[% Mujeres con discapacidad],">1")+COUNTIF(tblPerfiles[% Hombres con discapacidad],">1")+COUNTIF(tblPerfiles[% Niñas con discapacidad],">1")+COUNTIF(tblPerfiles[% Niños con discapacidad],">1")', "ROJO", "Es el % dentro de cada grupo: máximo 100%."),
    ("Actividades con beneficiarios pero sin perfil (no se pueden desagregar)",
     'SUMPRODUCT((tblWBS[Nivel]="Actividad")*(tblWBS[Perfil de beneficiarios]="")*(SUMIFS(tblWBS[Beneficiarios],tblWBS[Código WBS],tblWBS[Código WBS])>0))', "AMARILLO", "Asignar un perfil en la WBS (columna 'Perfil de beneficiarios')."),
    ("Perfil asignado que no existe en la hoja Perfiles",
     'SUMPRODUCT((tblWBS[Perfil de beneficiarios]<>"")*(COUNTIF(tblPerfiles[Perfil],tblWBS[Perfil de beneficiarios])=0))', "ROJO", "Elegir el perfil de la lista."),
    ("Perfiles en uso no confirmados",
     'SUMPRODUCT((tblPerfiles[Perfil]<>"")*(COUNTIF(tblWBS[Perfil de beneficiarios],tblPerfiles[Perfil])>0)*(tblPerfiles[Estado]<>"Confirmado"))', "AMARILLO", "Validar los % con el equipo técnico o con datos de proyectos anteriores."),
    ("Lugares sin departamento o sin zona", 'SUMPRODUCT((tblLugares[Lugar]<>"")*((tblLugares[Departamento]="")+(tblLugares[Zona]="")>0))', "ROJO", "Completar la hoja Lugares."),
    ("TARIFAS Y CARGOS", None, None, None),
    ("Tarifas sin fuente", 'COUNTIFS(tblTarifas[ID tarifa],"<>",tblTarifas[Fuente],"")', "ROJO", "Toda tarifa necesita cotización, factura o política."),
    ("Tarifas duplicadas (mismo ID y zona: el libro las SUMARÍA)", 'SUMPRODUCT((tblTarifas[ID tarifa]<>"")*(COUNTIFS(tblTarifas[ID tarifa],tblTarifas[ID tarifa],tblTarifas[Zona],tblTarifas[Zona])>1))', "ROJO", "Dejar una sola fila por ID + zona."),
    ("Tarifas sin moneda (o TRM vacía)", 'IF(TRM="",1,0)+COUNTIFS(tblTarifas[ID tarifa],"<>",tblTarifas[Moneda],"")', "ROJO", "Completar la moneda de cada tarifa y la TRM."),
    ("Tarifas usadas en este proyecto que no están confirmadas",
'COUNTIFS(tblTarifas[En uso],"Sí",tblTarifas[Estado],"<>Confirmado")',
     "AMARILLO", "Actualizar con cotización vigente y marcar Confirmado."),
    ("Tarifas en uso sin fecha o con más de 12 meses", 'COUNTIFS(tblTarifas[En uso],"Sí",tblTarifas[Fecha],"")+COUNTIFS(tblTarifas[En uso],"Sí",tblTarifas[Fecha],"<"&EDATE(TODAY(),-12))', "AMARILLO", "Actualizar el precio o confirmar que sigue vigente, con fecha (filtre Tarifas por 'En uso' = Sí)."),
    ("Cargos sin fuente", 'COUNTIFS(tblCargos[Cargo],"<>",tblCargos[Fuente],"")', "ROJO", "Todo costo de personal necesita fuente: tabla salarial, contrato o planilla."),
    ("Cargos usados con costo no confirmado", 'SUMPRODUCT((tblCargos[Cargo]<>"")*(tblCargos[Estado]<>"Confirmado")*(COUNTIF(tblPersonal[Cargo],tblCargos[Cargo])>0))', "AMARILLO", "Validar con Talento Humano."),
    ("PERSONAL Y OTROS COSTOS", None, None, None),
    ("Personal: cargo sin costo en Cargos", 'SUMPRODUCT((tblPersonal[Cargo]<>"")*(tblPersonal[Costo mensual COP]=0))', "ROJO", "Agregar el cargo en Cargos o corregir el nombre."),
    ("Personal: dedicación mayor a 100%", 'COUNTIF(tblPersonal[Dedicación],">1")', "ROJO", "Escribir la dedicación como porcentaje (15%)."),
    ("Personal: código WBS inexistente", 'SUMPRODUCT((tblPersonal[Cargo]<>"")*(COUNTIF(tblWBS[Código WBS],tblPersonal[Código WBS])=0))', "ROJO", "Corregir el código."),
    ("Otros costos: tarifa no encontrada", 'SUMPRODUCT((tblOtros[ID tarifa]<>"")*(tblOtros[Tarifa COP]=0))', "ROJO", "Crear la tarifa o corregir el ID."),
    ("Otros costos: código WBS inexistente o sin bloque",
     'SUMPRODUCT((tblOtros[ID tarifa]<>"")*(COUNTIF(tblWBS[Código WBS],tblOtros[Código WBS])=0))+SUMPRODUCT((tblOtros[ID tarifa]<>"")*(COUNTIF(tblMapeo[Bloque],tblOtros[Bloque])=0))', "ROJO", "Completar código; si falta el bloque, asignarlo en la tarifa (hoja Tarifas)."),
    ("DONANTE", None, None, None),
    ("% de indirectos por encima del máximo del donante", "IF(PctInd>MaxInd,1,0)", "ROJO", "Ajustar el % en Parametros."),
]
r = 7; n = 0
for lab, expr, sev, todo in checks:
    if expr is None:
        x = wsA.cell(row=r, column=2, value=lab); x.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
        for c in range(1, 6): wsA.cell(row=r, column=c).fill = PatternFill("solid", fgColor="5B7F8A")
        r += 1; continue
    n += 1
    wsA.cell(row=r, column=1, value=n); wsA.cell(row=r, column=2, value=lab)
    wsA.cell(row=r, column=3, value="=" + expr)
    wsA.cell(row=r, column=4, value=f'=IF(C{r}=0,"VERDE","{sev}")')
    wsA.cell(row=r, column=5, value=todo)
    for c in range(1, 6):
        x = wsA.cell(row=r, column=c); x.font = BF; x.border = BRD; x.fill = CAL
        x.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c == 4 else None)
    wsA.cell(row=r, column=3).number_format = '#,##0.##'
    wsA.cell(row=r, column=4).font = Font(name=FONT, size=9, bold=True)
    r += 1
extra = [
    ("Techo del donante: excedente (moneda de presentación)", '=IF(Techo="","",MAX(0,Res_TotalCOP*FX-Techo))', '=IF(Techo="","SIN DATO",IF(C{r}=0,"VERDE","ROJO"))',
     "Recortar vía supuestos (menos eventos, días, personas), nunca en el formato del donante."),
    ("Formato del donante vs. este modelo: diferencia (moneda de presentación)", '=IF(TotalDonante="","",ROUND(TotalDonante-Res_TotalCOP*FX,2))',
     '=IF(TotalDonante="","SIN DATO",IF(ABS(C{r})<1,"VERDE","ROJO"))', "Pegar en Parametros el total del Excel del donante. Si no cuadra, alguien editó el formato a mano."),
]
for lab, f1, f2, todo in extra:
    n += 1
    wsA.cell(row=r, column=1, value=n); wsA.cell(row=r, column=2, value=lab)
    wsA.cell(row=r, column=3, value=f1); wsA.cell(row=r, column=4, value=f2.format(r=r)); wsA.cell(row=r, column=5, value=todo)
    for c in range(1, 6):
        x = wsA.cell(row=r, column=c); x.font = BF; x.border = BRD; x.fill = CAL
        x.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c == 4 else None)
    wsA.cell(row=r, column=4).font = Font(name=FONT, size=9, bold=True)
    wsA.cell(row=r, column=3).number_format = '#,##0.##'
    r += 1
AL = r - 1
rng = f"D7:D{AL}"
for txt, fill in [("ROJO", RED), ("AMARILLO", YEL), ("SIN DATO", YEL), ("VERDE", GRN)]:
    wsA.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{txt}"'], fill=fill))
name("Aud_Estados", f"Auditoria!$D$7:$D${AL}")
wsA["A4"] = f'=COUNTIF(Aud_Estados,"ROJO")&" rojos · "&COUNTIF(Aud_Estados,"AMARILLO")&" amarillos · "&COUNTIF(Aud_Estados,"VERDE")&" verdes     |  Total proyecto: "&{mtxt("Res_TotalCOP*FX")}'
wsA["A4"].font = Font(name=FONT, size=12, bold=True)
wsA.freeze_panes = "A7"

# ================================================================= TABLERO (new)
wsB = wb.create_sheet("Tablero")
title(wsB, "Tablero – el proyecto de un vistazo", "Solo lectura. Todo se calcula desde WBS, Eventos, Personal y OtrosCostos. Moneda: la de presentación (Parametros).")
wsB.column_dimensions["A"].width = 9; wsB.column_dimensions["B"].width = 40
for c in range(3, 26):
    wsB.column_dimensions[L(c)].width = 16 if c <= 9 else 12.5
# header info
info = [("Proyecto", "=Proyecto"), ("Donante", "=DonanteActivo"), ("Moneda", "=Moneda"), ("TRM", "=TRM")]
for i, (lab, f) in enumerate(info):
    c = 1 + i * 3 if i else 1
for (cl, lab, cv, f) in [("A", "Proyecto:", "B", "=Proyecto"), ("C", "Donante:", "D", "=DonanteActivo"), ("E", "Moneda:", "F", "=Moneda"), ("G", "TRM:", "H", "=TRM")]:
    wsB[f"{cl}4"] = lab; wsB[f"{cl}4"].font = BOLD
    wsB[f"{cv}4"] = f; wsB[f"{cv}4"].font = BF
wsB["H4"].number_format = COP
wsB["A5"] = "Auditoría:"; wsB["A5"].font = BOLD
wsB["B5"] = '=COUNTIF(Aud_Estados,"ROJO")&" rojos · "&COUNTIF(Aud_Estados,"AMARILLO")&" amarillos  →  ver hoja Auditoria"'
wsB["B5"].font = BOLD
wsB.conditional_formatting.add("B5", FormulaRule(formula=['COUNTIF(Aud_Estados,"ROJO")>0'], fill=RED))
wsB.conditional_formatting.add("B5", FormulaRule(formula=['COUNTIF(Aud_Estados,"ROJO")=0'], fill=GRN))
# KPI tiles
tiles1 = [
    ("Total proyecto", "=Res_TotalCOP*FX", PRES, True), ("Costos directos", "=Res_DirectoCOP*FX", PRES, True), ("Indirectos", "=Res_IndCOP*FX", PRES, True),
    ("Actividades", '=COUNTIF(tblWBS[Nivel],"Actividad")', NUM, False), ("Eventos", "=SUM(tblEventos[Eventos contados])", NUM, False),
    ("Departamentos", "=SUMPRODUCT(--(COUNTIF(tblEventos[Departamento],Lista_Departamentos)>0))", NUM, False),
    ("Lugares", '=SUMPRODUCT((tblEventos[Lugar]<>"")/COUNTIF(tblEventos[Lugar],tblEventos[Lugar]&""))', NUM, False),
    ("Personas en el equipo", "=SUM(tblPersonal[Número de personas])", NUM, False),
]
tiles2 = [
    ("Participantes únicos", "=SUM(tblEventos[Participantes únicos])", NUM, False),
    ("Asistencias (part. + acomp.)", "=SUMPRODUCT((tblEventos[Participantes usados]+tblEventos[Acompañantes])*tblEventos[Repeticiones usadas])", NUM, False),
    ("Noches de hotel staff", "=SUM(tblEventos[Q Hospedaje staff])", NUM, False),
    ("Hab.-noches participantes", "=SUM(tblEventos[Q Hospedaje participantes])+SUMPRODUCT((tblOtros[ID tarifa]=\"HOSP-PART\")*tblOtros[Cantidad]*tblOtros[Factor])", NUM, False),
    ("Tiquetes aéreos", "=SUM(tblEventos[Q Tiquetes aéreos staff])+SUMPRODUCT((tblOtros[ID tarifa]=\"TIQ\")*tblOtros[Cantidad]*tblOtros[Factor])", NUM, False),
    ("Meses-persona", "=SUMPRODUCT(tblPersonal[Número de personas]*tblPersonal[Meses]*tblPersonal[Dedicación])", '#,##0.0', False),
    ("Dispositivos", '=SUMPRODUCT((tblOtros[Bloque]="Dispositivos de movilidad")*tblOtros[Cantidad]*tblOtros[Factor])', NUM, False),
    ("Supuestos por confirmar", '=COUNTIF(tblEventos[Estado],"Por confirmar")+COUNTIF(tblPersonal[Estado],"Por confirmar")+COUNTIF(tblOtros[Estado],"Por confirmar")', NUM, False),
]
big = Font(name=FONT, size=12, bold=True, color="1F4E5A")
_tot = 'SUMIFS(tblWBS[{c}],tblWBS[Suma al total],"<>No")'
tiles3 = [
    ("Beneficiarios (total)", "=" + _tot.format(c="Beneficiarios"), NUM, False),
    ("Mujeres", "=" + _tot.format(c="Mujeres"), NUM, False), ("Hombres", "=" + _tot.format(c="Hombres"), NUM, False),
    ("Niñas", "=" + _tot.format(c="Niñas"), NUM, False), ("Niños", "=" + _tot.format(c="Niños"), NUM, False),
    ("Con discapacidad", "=" + _tot.format(c="Total con discapacidad"), NUM, False),
    ("% mujeres y niñas", "=IF(B14=0,0,(C14+E14)/B14)", "0%", False),
    ("Sin desagregar (sin perfil)", "=B14-SUM(C14:F14)", NUM, False),
]
for rowl, tiles in [(7, tiles1), (10, tiles2), (13, tiles3)]:
    for i, (lab, f, fmt, money) in enumerate(tiles):
        c = 2 + i
        x = wsB.cell(row=rowl, column=c, value=lab); x.font = Font(name=FONT, size=8, bold=True, color="595959"); x.fill = TILE
        x.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        y = wsB.cell(row=rowl + 1, column=c, value=f); y.font = big; y.fill = TILE; y.number_format = fmt
        y.alignment = Alignment(horizontal="center", vertical="center")
        if money: money_cf(wsB, y.coordinate)
    wsB.row_dimensions[rowl].height = 26; wsB.row_dimensions[rowl + 1].height = 26
wsB.column_dimensions["B"].width = 40
# Department section
DR = 17
wsB.cell(row=DR - 1, column=1, value="Por departamento (solo eventos; personal y otros costos no tienen lugar)").font = H2
for i, h in enumerate(["#", "Departamento", "Eventos", "Lugares", "Participantes únicos", "Noches staff", "Costo eventos", "% de eventos"], 1):
    hcell(wsB, DR, i, h)
DEP_K = lambda k: (f'IFERROR(INDEX(Lista_Departamentos,SMALL(IF(COUNTIF(tblEventos[Departamento],Lista_Departamentos)>0,'
                   f'ROW(Lista_Departamentos)-MIN(ROW(Lista_Departamentos))+1),{k})),"")')
for k in range(1, MAX_DEP + 1):
    r = DR + k
    bcell(wsB, r, 1, k)
    bcell(wsB, r, 2, None); wsB[f"B{r}"] = ArrayFormula(f"B{r}", "=" + DEP_K(k))
    d = f"$B{r}"
    bcell(wsB, r, 3, f'=IF({d}="","",SUMIFS(tblEventos[Eventos contados],tblEventos[Departamento],{d}))', NUM)
    bcell(wsB, r, 4, f'=IF({d}="","",SUMPRODUCT((tblEventos[Departamento]={d})*(tblEventos[Lugar]<>"")/COUNTIFS(tblEventos[Departamento],tblEventos[Departamento]&"",tblEventos[Lugar],tblEventos[Lugar]&"")))', NUM)
    bcell(wsB, r, 5, f'=IF({d}="","",SUMIFS(tblEventos[Participantes únicos],tblEventos[Departamento],{d}))', NUM)
    bcell(wsB, r, 6, f'=IF({d}="","",SUMIFS(tblEventos[Q Hospedaje staff],tblEventos[Departamento],{d}))', NUM)
    bcell(wsB, r, 7, f'=IF({d}="","",SUMIFS(tblEventos[Total COP],tblEventos[Departamento],{d})*FX)', PRES)
    bcell(wsB, r, 8, f'=IF({d}="","",IF(SUM(tblEventos[Total COP])=0,0,SUMIFS(tblEventos[Total COP],tblEventos[Departamento],{d})/SUM(tblEventos[Total COP])))', "0.0%")
money_cf(wsB, f"G{DR+1}:G{DR+MAX_DEP}")
MC = 10
wsB.cell(row=DR - 1, column=MC, value="Beneficiarios: sexo, edad y discapacidad (estimado)").font = H2
for i, h in enumerate(["Grupo", "Total", "Con discapacidad", "Sin discapacidad", "% del total"]):
    hcell(wsB, DR, MC + i, h)
for i, g in enumerate(G4 + ["Total"]):
    r = DR + 1 + i
    fnt = BOLD if g == "Total" else BF
    bcell(wsB, r, MC, g, font=fnt)
    if g == "Total":
        bcell(wsB, r, MC + 1, f"=SUM({L(MC+1)}{DR+1}:{L(MC+1)}{DR+4})", NUM, font=fnt)
        bcell(wsB, r, MC + 2, f"=SUM({L(MC+2)}{DR+1}:{L(MC+2)}{DR+4})", NUM, font=fnt)
    else:
        bcell(wsB, r, MC + 1, "=" + _tot.format(c=g), NUM, font=fnt)
        bcell(wsB, r, MC + 2, "=" + _tot.format(c=f"{g} con discapacidad"), NUM, font=fnt)
    bcell(wsB, r, MC + 3, f"={L(MC+1)}{r}-{L(MC+2)}{r}", NUM, font=fnt)
    bcell(wsB, r, MC + 4, f"=IF($B$14=0,0,{L(MC+1)}{r}/$B$14)", "0.0%", font=fnt)
bcell(wsB, DR + 6, MC, "Sin desagregar", font=SUB); bcell(wsB, DR + 6, MC + 1, "=I14", NUM, font=SUB)
wsB.cell(row=DR + 7, column=MC, value="Solo actividades con 'Suma al total' ≠ No. Acompañantes no cuentan.").font = SUB
# Activity section
AR = DR + MAX_DEP + 4
wsB.cell(row=AR - 2, column=1, value="Actividades: metas vs. presupuesto y distribución territorial").font = H2
wsB.cell(row=AR - 1, column=1, value="Las columnas de la derecha muestran cuántos eventos tiene cada actividad en cada departamento: un hueco donde debería haber eventos = algo falta.").font = SUB
acth = ["Código", "Actividad", "Plantilla", "Eventos declarados", "Eventos presup.", "Particip. declarados", "Particip. presup.", "Lugares", "Control", "Total actividad", "% del directo"]
for i, h in enumerate(acth, 1):
    hcell(wsB, AR, i, h)
DC0 = len(acth) + 1
for j in range(MAX_DEP):
    c = DC0 + j
    x = wsB.cell(row=AR, column=c, value=f'=IF($B${DR+1+j}="","",$B${DR+1+j})')
    x.font = HF; x.border = BRD
    x.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center", text_rotation=0)
wsB.row_dimensions[AR].height = 40
wsB.conditional_formatting.add(f"{L(DC0)}{AR}:{L(DC0+MAX_DEP-1)}{AR}", FormulaRule(formula=[f'{L(DC0)}{AR}<>""'], fill=PatternFill("solid", fgColor="2E75B6")))
for k in range(1, MAX_ACT_TAB + 1):
    r = AR + k
    bcell(wsB, r, 1, None, TXT); wsB[f"A{r}"] = ArrayFormula(f"A{r}", "=" + ACT_K(k))
    a = f"$A{r}"
    look = lambda col: f'INDEX(tblWBS[{col}],MATCH({a},tblWBS[Código WBS],0))'
    bcell(wsB, r, 2, f'=IF({a}="","",{look("Nombre")})')
    bcell(wsB, r, 3, f'=IF({a}="","",T({look("Plantilla")}))')
    bcell(wsB, r, 4, f'=IF({a}="","",IF({look("Eventos declarados")}="","",{look("Eventos declarados")}))', NUM)
    bcell(wsB, r, 5, f'=IF({a}="","",{look("Eventos presupuestados")})', NUM)
    bcell(wsB, r, 6, f'=IF({a}="","",IF({look("Participantes declarados")}="","",{look("Participantes declarados")}))', NUM)
    bcell(wsB, r, 7, f'=IF({a}="","",{look("Participantes presupuestados")})', NUM)
    bcell(wsB, r, 8, f'=IF({a}="","",SUMPRODUCT((tblEventos[Código WBS]={a})*(tblEventos[Lugar]<>"")/COUNTIFS(tblEventos[Código WBS],tblEventos[Código WBS]&"",tblEventos[Lugar],tblEventos[Lugar]&"")))', NUM)
    bcell(wsB, r, 9, f'=IF({a}="","",{look("Control")})')
    bcell(wsB, r, 10, f'=IF({a}="","",{look("Total")})', PRES)
    bcell(wsB, r, 11, f'=IF({a}="","",IF(Res_DirectoCOP=0,0,{look("Total")}/(Res_DirectoCOP*FX)))', "0.0%")
    for j in range(MAX_DEP):
        c = DC0 + j; hd = f"{L(c)}${AR}"
        bcell(wsB, r, c, f'=IF(OR({a}="",{hd}=""),"",SUMIFS(tblEventos[Eventos contados],tblEventos[Código WBS],{a},tblEventos[Departamento],{hd}))', NUM)
        wsB.cell(row=r, column=c).alignment = Alignment(horizontal="center")
AT = AR + MAX_ACT_TAB + 1
wsB.cell(row=AT, column=2, value="TOTAL").font = BOLD
for c, fmt in [(4, NUM), (5, NUM), (6, NUM), (7, NUM), (10, PRES)]:
    x = bcell(wsB, AT, c, f"=SUM({L(c)}{AR+1}:{L(c)}{AR+MAX_ACT_TAB})", fmt, font=BOLD)
for j in range(MAX_DEP):
    c = DC0 + j
    bcell(wsB, AT, c, f'=IF({L(c)}${AR}="","",SUM({L(c)}{AR+1}:{L(c)}{AR+MAX_ACT_TAB}))', NUM, font=BOLD)
x = bcell(wsB, AT + 1, 2, "Personal + otros costos sin actividad de eventos ya están incluidos en 'Total actividad'.", font=SUB, fill=PatternFill(fill_type=None))
ctlr = f"I{AR+1}:I{AR+MAX_ACT_TAB}"
wsB.conditional_formatting.add(ctlr, CellIsRule(operator="equal", formula=['"REVISAR"'], fill=RED))
wsB.conditional_formatting.add(ctlr, CellIsRule(operator="equal", formula=['"SIN METAS"'], fill=YEL))
wsB.conditional_formatting.add(ctlr, CellIsRule(operator="equal", formula=['"OK"'], fill=GRN))
mat = f"{L(DC0)}{AR+1}:{L(DC0+MAX_DEP-1)}{AR+MAX_ACT_TAB}"
wsB.conditional_formatting.add(mat, CellIsRule(operator="greaterThan", formula=['0'], fill=PatternFill("solid", fgColor="BDD7EE")))
money_cf(wsB, f"J{AR+1}:J{AT}")
BR = AT + 4
wsB.cell(row=BR - 2, column=1, value="Beneficiarios por actividad (estimados con el perfil de cada actividad)").font = H2
bh = ["Código", "Actividad", "Perfil", "Beneficiarios", "Suma al total", "Mujeres", "Hombres", "Niñas", "Niños",
      "Mujeres c/disc.", "Hombres c/disc.", "Niñas c/disc.", "Niños c/disc.", "Total c/disc."]
for i, h in enumerate(bh, 1):
    hcell(wsB, BR, i, h)
wsB.row_dimensions[BR].height = 32
bcols = ["Beneficiarios", "Suma al total"] + G4 + [f"{g} con discapacidad" for g in G4] + ["Total con discapacidad"]
for k in range(1, MAX_ACT_TAB + 1):
    r = BR + k; a = f"$A{r}"
    bcell(wsB, r, 1, f"=A{AR+k}", TXT)
    look = lambda col: f'INDEX(tblWBS[{col}],MATCH({a},tblWBS[Código WBS],0))'
    bcell(wsB, r, 2, f'=IF({a}="","",{look("Nombre")})')
    bcell(wsB, r, 3, f'=IF({a}="","",T({look("Perfil de beneficiarios")}))')
    for j, col in enumerate(bcols):
        c = 4 + j
        if col == "Suma al total":
            bcell(wsB, r, c, f'=IF({a}="","",IF(T({look(col)})="No","No","Sí"))')
            wsB.cell(row=r, column=c).alignment = Alignment(horizontal="center")
        else:
            bcell(wsB, r, c, f'=IF({a}="","",{look(col)})', NUM)
BT = BR + MAX_ACT_TAB + 1
bcell(wsB, BT, 2, "TOTAL PROYECTO (solo 'Suma al total' = Sí)", font=BOLD)
for j, col in enumerate(bcols):
    if col == "Suma al total": continue
    bcell(wsB, BT, 4 + j, "=" + _tot.format(c=col), NUM, font=BOLD)
wsB.conditional_formatting.add(f"C{BR+1}:C{BR+MAX_ACT_TAB}", FormulaRule(formula=[f'AND($A{BR+1}<>"",$C{BR+1}="",N($D{BR+1})>0)'], fill=YEL))
wsB.conditional_formatting.add(f"E{BR+1}:E{BR+MAX_ACT_TAB}", CellIsRule(operator="equal", formula=['"No"'], font=Font(color="C00000", bold=True)))
wsB.freeze_panes = "C4"
wsB.page_setup.orientation = "landscape"; wsB.page_setup.fitToWidth = 1; wsB.page_setup.fitToHeight = 0
wsB.sheet_properties.pageSetUpPr.fitToPage = True
wsB.print_area = f"A1:{L(DC0+MAX_DEP-1)}{BT}"

# ================================================================= CONCILIACION (only for reconstructions)
CON = PROJ.get("conciliacion")
if CON:
    wsX = wb.create_sheet("Conciliacion")
    title(wsX, CON["titulo"], CON["subtitulo"])
    heads = [("Línea aprobada", 10), ("Concepto (presupuesto aprobado)", 52), ("Aprobado USD", 14), ("Dónde está en el modelo", 46), ("Modelo USD", 14), ("Diferencia USD", 13), ("Comentario", 70)]
    for i, (h, w) in enumerate(heads, 1):
        hcell(wsX, 4, i, h); wsX.column_dimensions[L(i)].width = w
    r = 5
    for it in CON["lineas"]:
        if "seccion" in it:
            x = wsX.cell(row=r, column=1, value=it["seccion"]); x.font = Font(name=FONT, size=9, bold=True, color="FFFFFF")
            for c in range(1, 8): wsX.cell(row=r, column=c).fill = PatternFill("solid", fgColor="5B7F8A")
            r += 1; continue
        bcell(wsX, r, 1, it["linea"]); bcell(wsX, r, 2, it["concepto"]); bcell(wsX, r, 3, it["aprobado"], USD)
        bcell(wsX, r, 4, it["donde"]); bcell(wsX, r, 5, "=" + it["expr"], USD)
        bcell(wsX, r, 6, f"=IF(ABS(E{r}-C{r})<0.01,0,ROUND(E{r}-C{r},2))", USD); bcell(wsX, r, 7, it.get("nota"))
        for c in (2, 4, 7): wsX.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    last = r - 1
    wsX.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=["ABS(F5)>=0.01"], fill=RED))
    wsX.conditional_formatting.add(f"F5:F{last}", FormulaRule(formula=['AND(ISNUMBER(F5),ABS(F5)<0.01)'], fill=GRN))
    r += 1
    wsX.cell(row=r, column=1, value="Totales").font = H2; r += 1
    for lab, apr, f in [("Total costos directos", CON["directo"], "=Res_DirectoCOP/TRM"), ("Costos indirectos (PSC)", CON["psc"], "=Res_IndCOP/TRM"),
                        ("TOTAL PROYECTO", CON["total"], "=Res_TotalCOP/TRM")]:
        bcell(wsX, r, 2, lab, font=BOLD); bcell(wsX, r, 3, apr, USD, font=BOLD); bcell(wsX, r, 5, f, USD, font=BOLD)
        bcell(wsX, r, 6, f"=ROUND(E{r}-C{r},2)", USD, font=BOLD)
        wsX.conditional_formatting.add(f"F{r}", FormulaRule(formula=[f"ABS(F{r})>=0.01"], fill=RED))
        wsX.conditional_formatting.add(f"F{r}", FormulaRule(formula=[f"ABS(F{r})<0.01"], fill=GRN))
        r += 1
    r += 1
    wsX.cell(row=r, column=1, value="Por categoría del donante").font = H2; r += 1
    for lab, apr in CON["categorias"]:
        bcell(wsX, r, 2, lab); bcell(wsX, r, 3, apr, USD)
        bcell(wsX, r, 5, f'=SUMIF(Res_CatNom,B{r},Res_CatVal)/TRM', USD); bcell(wsX, r, 6, f"=ROUND(E{r}-C{r},2)", USD)
        wsX.conditional_formatting.add(f"F{r}", FormulaRule(formula=[f"ABS(F{r})>=0.01"], fill=YEL))
        wsX.conditional_formatting.add(f"F{r}", FormulaRule(formula=[f"ABS(F{r})<0.01"], fill=GRN))
        r += 1
    r += 1
    wsX.cell(row=r, column=1, value="Hallazgos al pasar el presupuesto aprobado al modelo").font = H2; r += 1
    for h in CON["hallazgos"]:
        x = wsX.cell(row=r, column=2, value=h); x.font = BF; x.alignment = Alignment(wrap_text=True, vertical="top")
        wsX.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7); wsX.row_dimensions[r].height = 30
        r += 1
    wsX.freeze_panes = "C5"
    wsX.sheet_properties.tabColor = "7030A0"

# ================================================================= LEEME
wsI = wb.create_sheet("Leeme")
title(wsI, "Constructor de presupuestos CIREC – v4.3", "Guía rápida. Datos de EJEMPLO (se pueden borrar: solo celdas amarillas).")
wsI.column_dimensions["A"].width = 3; wsI.column_dimensions["B"].width = 130
txt = [
    ("REGLA DE ORO", True),
    ("Nunca se escribe un total a mano. Se escriben supuestos (qué actividad, dónde, cuántas veces) y el libro calcula.", False),
    ("", False),
    ("PROYECTO NUEVO EN 6 PASOS (orden de las pestañas azules)", True),
    ("1. Parametros: nombre, donante, moneda, TRM, % indirectos, techo.", False),
    ("2. WBS: pegar el marco lógico (código, nivel, nombre). Por cada actividad: plantilla de evento + nº de eventos y participantes que DICE la propuesta.", False),
    ("3. Eventos: una fila por cada LUGAR donde ocurre cada actividad. Mínimo: Código WBS + Lugar. Repeticiones si se repite en el mismo lugar (fases).", False),
    ("4. Personal: cargo × nº × meses × dedicación.   5. OtrosCostos: dispositivos, kits masivos, seguros, equipos (Código + ID tarifa + cantidad).", False),
    ("6. Tablero: ¿cuadran eventos, lugares y participantes por actividad y departamento?  Auditoria: cero ROJOS antes del formato del donante.", False),
    ("", False),
    ("LO QUE SE LLENA UNA SOLA VEZ (pestañas amarillas = banco de CIREC)", True),
    ("Plantillas: el perfil típico de cada tipo de evento (staff, días, noches, participantes, comida, transporte, kit). Cambiar una plantilla cambia todos sus eventos.", False),
    ("Perfiles: % de mujeres, hombres, niñas y niños de cada tipo de actividad, y % con discapacidad dentro de cada grupo. Cada actividad de la WBS elige su perfil.", False),
    ("Lugares: municipio → departamento + zona tarifaria. Tarifas: precios con fuente, fecha y bloque. Cargos: costo mensual por cargo. MapeoDonante: bloque → categoría de cada donante.", False),
    ("", False),
    ("CUANDO UN EVENTO ES DISTINTO A SU PLANTILLA", True),
    ("Participantes: columna visible 'Participantes (si difiere)'. Todo lo demás: abra el grupo [+] encima de las columnas y escriba solo en la columna 'Ajuste …' que cambia.", False),
    ("Una fila con ajustes queda marcada en naranja en la columna ID. Vacío = valor de la plantilla; 0 = cero.", False),
    ("", False),
    ("ATAJOS PARA IR RÁPIDO", True),
    ("Varias filas iguales: llene una, selecciónela con las de abajo y Ctrl+D (copia hacia abajo). Luego cambie solo el Lugar.", False),
    ("Pegar desde otro Excel: copie columnas Código y Lugar y péguelas en la primera fila vacía debajo de la tabla (Pegado especial → Valores). La tabla crece sola.", False),
    ("Listas desplegables: en Excel 365/online puede escribir las primeras letras para filtrar. En Excel 2019 use Alt+↓ para abrir la lista.", False),
    ("Borrar filas: clic derecho dentro de la tabla → Eliminar → Filas de la tabla. Nunca escriba sobre celdas grises: la Auditoría lo detecta.", False),
    ("", False),
    ("EVENTO vs ACTIVIDAD", True),
    ("Actividad = QUÉ hace el proyecto (WBS, tiene metas). Evento = DÓNDE y CUÁNTAS VECES ocurre. Un lugar distinto = una fila distinta (nunca multiplique lugares).", False),
    ("Repeticiones = el mismo evento, en el mismo lugar, otra vez (fases de brigada, talleres mensuales). 'Mismos participantes' en la plantilla decide si se cuentan una vez o cada vez.", False),
    ("Talleres masivos e idénticos (ej. 750 talleres ERAE): NO cree 750 filas. Una fila por departamento o municipio, Repeticiones = nº de talleres, y una plantilla con 'Paquete por evento' (precio todo incluido por taller). El costo, el nº de eventos y los beneficiarios salen de esas mismas filas.", False),
    ("Complementaria = Sí: segundo grupo de un mismo evento con logística distinta (ej. 2 personas más que van por tierra). No cuenta como evento.", False),
    ("", False),
    ("BENEFICIARIOS", True),
    ("Por defecto, los beneficiarios de una actividad = participantes únicos de sus eventos (sin acompañantes). Si la actividad no tiene eventos (talleres masivos, transferencias), escriba el número en WBS → 'Beneficiarios (si difiere)'.", False),
    ("El reparto por sexo y edad redondea cada grupo y ajusta la diferencia en el grupo mayor: la suma siempre da el total. Si las mismas personas aparecen en dos actividades, marque 'Suma al total' = No en una de ellas.", False),
    ("", False),
    ("CONVENCIONES", True),
    ("Días transporte staff = días con transporte local (aeropuerto + internos); puede ser menor que los días en sitio.", False),
    ("Días en sitio = días contando los de viaje. Noches = noches que duermen fuera (normalmente días − 1). Códigos WBS en texto (1.10 ≠ 1.1). Dedicación como %.", False),
    ("Moneda: todo se calcula en COP. 'Moneda de presentación' solo cambia cómo se ven totales y narrativas. Cada tarifa puede cotizarse en COP o USD.", False),
    ("Indirectos: en Parametros se puede calcular sobre costos directos o sobre el total del proyecto.", False),
    ("", False),
    ("BANCO DE PRECIOS", True),
    ("Tarifas y Cargos son el banco de precios de CIREC: cada fila lleva fuente, fecha y estado. Mantenga una copia maestra 'plantilla' y copie el archivo para cada proyecto.", False),
    ("No cambie estos 7 ID de Tarifas (Eventos los usa): HOSP-STAFF, HOSP-PART, ALIM, REFRIG, TIQ, TTE-STAFF, SALON.", False),
    ("Prótesis y órtesis: ID ORT-, PRO-, REP-PRO-, VAR- (lista de precios P&O 2025). Usar en OtrosCostos con su cantidad.", False),
    ("", False),
    ("POWER BI", True),
    ("Abra y guarde el archivo en Excel antes de actualizar Power BI. Power BI lee solo columnas COP. Ver Guia_PowerBI.md.", False),
]
for i, (t, b) in enumerate(txt):
    c = wsI.cell(row=4 + i, column=2, value=t)
    c.font = Font(name=FONT, size=10, bold=b, color="1F4E5A" if b else "000000")
    c.alignment = Alignment(wrap_text=True)

order = ["Leeme", "Tablero", "Parametros", "WBS", "Eventos", "Personal", "OtrosCostos", "Auditoria", "Resumen", "LineasDonante",
         "Plantillas", "Perfiles", "Lugares", "Tarifas", "Cargos", "MapeoDonante", "Listas"]
if CON: order.insert(2, "Conciliacion")
wb._sheets = [wb[s] for s in order]
colors = {"Leeme": "1F4E5A", "Tablero": "548235", "Parametros": "2E75B6", "WBS": "2E75B6", "Eventos": "2E75B6", "Personal": "2E75B6",
          "OtrosCostos": "2E75B6", "Auditoria": "C00000", "Resumen": "548235", "LineasDonante": "548235", "Plantillas": "BF8F00", "Perfiles": "BF8F00",
          "Lugares": "BF8F00", "Tarifas": "BF8F00", "Cargos": "BF8F00", "MapeoDonante": "BF8F00", "Listas": "7F7F7F"}
for s, c in colors.items():
    wb[s].sheet_properties.tabColor = c
wb["Listas"].sheet_state = "hidden"
wb.active = 1
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT, "eventos cols", len(ecols), "visible inputs")
json.dump({"keep": ["ID evento", "Código WBS", "Plantilla usada", "Lugar", "Departamento", "Zona", "Mes", "Año", "Estado"],
           "cop": [f"COP {b[0]}" for b in blocks]}, open(OUT + ".pq.json", "w"), ensure_ascii=False)
